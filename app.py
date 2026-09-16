"""DETRAN/MA — Sistema de Gestão de Exames Práticos.

Arquivo único, pronto para substituir o app.py do repositório.

Execução:
    pip install -r requirements.txt
    streamlit run app.py

Variáveis de ambiente opcionais:
    DETRAN_DB         caminho do banco SQLite (padrão: dados/detran.db)
    DETRAN_SNAPSHOTS  pasta dos snapshots JSON (padrão: dados/snapshots)
    DETRAN_LOG_LEVEL  nível de log (padrão: INFO)

Organização do arquivo:
    1. Configuração e constantes
    2. Regras de negócio (funções puras, sem Streamlit)
    3. Persistência (SQLite, backup, snapshots)
    4. Exportadores (PDF e Excel)
    5. Componentes de interface
    6. Abas
    7. Aplicação
"""

from __future__ import annotations

import base64
import calendar
import datetime
import hashlib
import html
import inspect
import io
import json
import logging
import os
import re
import sqlite3
import threading
import unicodedata
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Sequence

import holidays
import openpyxl
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.styles import Alignment, Border, Font, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

try:  # Componente opcional de calendário interativo.
    from streamlit_calendar import calendar as componente_calendario

    TEM_COMPONENTE_CALENDARIO = True
except ImportError:
    TEM_COMPONENTE_CALENDARIO = False


# ===========================================================================
# 1. CONFIGURAÇÃO E CONSTANTES
# ===========================================================================

SEPARADOR_CHAVE = "||"
SCHEMA_VERSION = 6

MESES_LISTA = [
    "Janeiro",
    "Fevereiro",
    "Março",
    "Abril",
    "Maio",
    "Junho",
    "Julho",
    "Agosto",
    "Setembro",
    "Outubro",
    "Novembro",
    "Dezembro",
]

DIAS_SEMANA_OPCOES = [
    "Segunda",
    "Terça",
    "Quarta",
    "Quinta",
    "Sexta",
    "Sábado",
    "Domingo",
]
DIAS_UTEIS_PADRAO = DIAS_SEMANA_OPCOES[:5]
ANOS_A_FRENTE = 5


def anos_disponiveis(anos_a_frente: int = ANOS_A_FRENTE) -> list[int]:
    """Anos selecionáveis, sempre a partir do ano corrente."""
    ano_atual = datetime.date.today().year
    return list(range(ano_atual, ano_atual + anos_a_frente))


def indice_mes_atual() -> int:
    return datetime.date.today().month - 1


def numero_do_mes(nome_mes: str) -> int:
    return MESES_LISTA.index(nome_mes) + 1


BANCAS_PADRAO: dict[str, list[str]] = {
    "Banca São Luís": [
        "São Luís Pátio",
        "São Luís Castelinho",
        "São Luís Cohatrac",
        "São Luís Cidade Operária",
        "Paço do Lumiar",
        "São José de Ribamar",
        "Raposa",
        "Pinheiro",
        "São Bento",
        "Carutapera",
        "Turilândia",
        "Tutóia",
        "Barreirinhas",
        "Chapadinha",
        "Axixá",
        "Rosário",
        "Santa Rita",
        "Coroatá",
        "Itapecuru-mirim",
        "Codó",
    ],
    "Banca Imperatriz": [
        "Imperatriz",
        "Açailândia",
        "Estreito",
        "Grajaú",
        "Amarante do Maranhão",
    ],
    "Banca Timon": ["Timon - Pátio", "Matões", "Parnarama"],
    "Banca Caxias": [
        "Caxias - Pátio",
        "Codó - Regional",
        "São João do Sóter",
        "Aldeias Altas",
    ],
    "Banca Bacabal": [
        "Bacabal - Pátio",
        "Pedreiras",
        "Lago da Pedra",
        "Olho d'Água das Cunhãs",
    ],
    "Banca Santa Inês": [
        "Santa Inês - Pátio",
        "Viana",
        "Zé Doca",
        "Monção",
        "Pindaré-Mirim",
    ],
}

# Ordem de apresentação nos relatórios em PDF: primeiro as unidades da sede,
# depois a região metropolitana e, por último, os demais municípios por ordem
# da data do primeiro exame lançado.
ORDEM_PRIORITARIA_PDF: dict[str, list[str]] = {
    "Banca São Luís": [
        "São Luís Pátio",
        "São Luís Castelinho",
        "São Luís Cohatrac",
        "São Luís Cidade Operária",
        "Paço do Lumiar",
        "São José de Ribamar",
        "Raposa",
    ],
}

GRUPOS_PDF: dict[str, list[str]] = {
    "São Luís Pátio": "SEDE — SÃO LUÍS",
    "São Luís Castelinho": "SEDE — SÃO LUÍS",
    "São Luís Cohatrac": "SEDE — SÃO LUÍS",
    "São Luís Cidade Operária": "SEDE — SÃO LUÍS",
    "Paço do Lumiar": "REGIÃO METROPOLITANA",
    "São José de Ribamar": "REGIÃO METROPOLITANA",
    "Raposa": "REGIÃO METROPOLITANA",
}
GRUPO_INTERIOR = "DEMAIS MUNICÍPIOS"
GRUPO_SEDE_GENERICO = "SEDE DA BANCA"

HORARIOS_PADRAO = [
    "08:30",
    "09:30",
    "10:30",
    "11:30",
    "13:30",
    "14:30",
    "15:30",
    "16:30",
]

CAPACIDADE_BANCAS_PADRAO: dict[str, int] = {
    "Banca São Luís": 34,
    "Banca Imperatriz": 14,
    "Banca Bacabal": 5,
    "Banca Caxias": 4,
    "Banca Timon": 4,
    "Banca Santa Inês": 3,
}
LIMITE_FORA_SEDE_PADRAO: dict[str, int] = {
    "Banca São Luís": 16,
    "Banca Imperatriz": 10,
    "Banca Bacabal": 5,
    "Banca Caxias": 4,
    "Banca Timon": 4,
    "Banca Santa Inês": 3,
}

BANCAS_APOIO_CONJUNTO = ["Banca Caxias", "Banca Timon"]
BANCAS_COM_QUADRO_MATRIZ = ["Banca São Luís", "Banca Imperatriz"]
BANCAS_COM_DISPONIBILIDADE_SEMANAL = ["Banca São Luís", "Banca Imperatriz"]

# ---------------------------------------------------------------------------
# Ajuste 3 — deslocamento das bancas itinerantes
# ---------------------------------------------------------------------------
# Minutos de deslocamento rodoviário ESTIMADOS a partir da sede da banca fixa
# até cada município. São valores de referência (distância média por rodovia
# convertida em tempo), não uma consulta em tempo real ao Google Maps — o
# sistema não tem acesso à internet em produção. Cada valor pode ser corrigido
# a qualquer momento no cadastro da viagem (o campo "minutos de deslocamento"
# é editável e a correção fica salva para a próxima vez que a localidade for
# usada), o que também cobre os casos em que uma banca atende um município
# normalmente listado sob outra banca fixa.
HORA_PARTIDA_PADRAO = datetime.time(8, 0)
DURACAO_TURMA_MIN = 60
DURACAO_ALMOCO_MIN = 60
LIMITE_FIM_EXPEDIENTE = datetime.time(17, 0)

TEMPO_DESLOCAMENTO_PADRAO_MIN: dict[str, dict[str, int]] = {
    "Banca São Luís": {
        # Sede e Região Metropolitana: sem deslocamento a calcular.
        "São Luís Pátio": 0,
        "São Luís Castelinho": 0,
        "São Luís Cohatrac": 0,
        "São Luís Cidade Operária": 0,
        "Paço do Lumiar": 0,
        "São José de Ribamar": 0,
        "Raposa": 0,
        # Itinerantes — estimativa rodoviária a partir de São Luís.
        "Rosário": 60,
        "Santa Rita": 75,
        "Itapecuru-mirim": 90,
        "Icatu": 120,
        "Axixá": 120,
        "Vitória do Mearim": 120,
        "Arari": 120,
        "Pinheiro": 120,
        "São Bento": 90,
        "Viana": 150,
        "Cantanhede": 150,
        "Coroatá": 150,
        "Turilândia": 180,
        "Codó": 240,
        "Chapadinha": 210,
        "Barreirinhas": 210,
        "Brejo": 300,
        "Tutóia": 300,
        "Carutapera": 330,
    },
    "Banca Imperatriz": {
        "Imperatriz": 0,
        "João Lisboa": 20,
        "Governador Edison Lobão": 30,
        "Cidelândia": 40,
        "Campestre do Maranhão": 60,
        "Açailândia": 60,
        "Estreito": 60,
        "Porto Franco": 75,
        "Itinga do Maranhão": 90,
        "Amarante do Maranhão": 90,
        "Alto Alegre do Maranhão": 120,
        "Bom Jesus das Selvas": 120,
        "Buriticupu": 150,
        "Grajaú": 180,
        "Balsas": 240,
    },
    "Banca Timon": {
        "Timon - Pátio": 0,
        "Matões": 60,
        "Coelho Neto": 40,
        "Caxias - Pátio": 120,
        "Aldeias Altas": 140,
        "São João do Sóter": 160,
        "Buriti Bravo": 120,
        "Colinas": 150,
        "Passagem Franca": 150,
        "Pastos Bons": 180,
        "São João dos Patos": 180,
        "Barra do Corda": 240,
    },
    "Banca Caxias": {
        "Caxias - Pátio": 0,
        "Aldeias Altas": 20,
        "São João do Sóter": 40,
        "Coelho Neto": 40,
        "Buriti Bravo": 60,
        "Codó - Regional": 60,
    },
    "Banca Bacabal": {
        "Bacabal - Pátio": 0,
        "Trizidela do Vale": 20,
        "Pedreiras": 30,
        "Olho d'Água das Cunhãs": 30,
        "Lago da Pedra": 45,
        "Vitorino Freire": 60,
        "São Mateus do Maranhão": 60,
        "Presidente Dutra": 60,
        "Dom Pedro": 90,
    },
    "Banca Santa Inês": {
        "Santa Inês - Pátio": 0,
        "Pindaré-Mirim": 40,
        "Monção": 60,
        "Santa Luzia": 60,
        "Viana": 90,
        "Zé Doca": 90,
    },
}

COLS_CATEGORIA = ["Cat A", "Cat B", "Cat C", "Cat D", "Cat E"]
COLS_PCD = ["PCD A", "PCD B", "PCD C", "PCD D", "PCD E"]
COLS_VAGAS = COLS_CATEGORIA + COLS_PCD

STATUS_DISPONIVEL = "Disponível"
STATUS_INDISPONIVEL = "Indisponível"
STATUS_FERIADO = "Feriado / Sem Atendimento"
OPCOES_STATUS = [STATUS_DISPONIVEL, STATUS_INDISPONIVEL, STATUS_FERIADO]
STATUS_BLOQUEADOS = {STATUS_INDISPONIVEL, STATUS_FERIADO}

TURNO_INTEGRAL = "Dia inteiro"
TURNO_MANHA = "Manhã"
TURNO_TARDE = "Tarde"
TURNO_SEM_ATENDIMENTO = "Não atende"
TURNOS = [TURNO_INTEGRAL, TURNO_MANHA, TURNO_TARDE]
TURNOS_POR_DIA = [TURNO_INTEGRAL, TURNO_MANHA, TURNO_TARDE, TURNO_SEM_ATENDIMENTO]

UF_FERIADOS = "MA"

# Logística padrão atrelada a cada equipe itinerante. Aparece apenas no PDF da
# escala de viagens, conforme solicitado — não há campo para isso na interface.
PREPOSTOS_POR_EQUIPE = 1
VEICULOS_POR_EQUIPE = 1

COR_PRIMARIA = "#1A365D"
COR_SECUNDARIA = "#2B6CB0"
COR_CLARA = "#EDF2F7"
COR_NEUTRA = "#718096"
COR_ALERTA = "#E53E3E"
COR_VIAGEM = "#2F855A"

# Destaque das linhas já lançadas no editor de calendário (Ajuste 2).
# Destaque das linhas já lançadas no editor de calendário (Ajuste 2) e do
# alerta de deslocamento (Ajuste 3).
COR_LANCADO_BG = "#C6F6D5"
COR_LANCADO_TXT = "#22543D"
COR_INATIVO_BG = "#E2E8F0"
COR_INATIVO_TXT = "#4A5568"
COR_ALERTA_DESLOC_BG = "#FED7D7"
COR_ALERTA_DESLOC_TXT = "#822727"

CSS_CUSTOMIZADO = """
<style>
.main { background-color: #f8f9fa; }
.header-container {
    background: linear-gradient(90deg, #1A365D 0%, #2B6CB0 100%);
    padding: 20px;
    border-radius: 10px;
    color: white;
    margin-bottom: 25px;
    box-shadow: 0 4px 6px rgba(0,0,0,0.1);
}
.header-container h1 { color: white !important; margin: 0; font-size: 26px; font-weight: 700; }
.header-container p { color: #E2E8F0; margin: 5px 0 0 0; font-size: 14px; }
.faixa-banca {
    background-color: #2B579A;
    color: #FFFFFF;
    padding: 8px;
    font-weight: bold;
    text-align: center;
    border-radius: 4px;
    font-size: 14px;
    text-transform: uppercase;
    margin-top: 15px;
    margin-bottom: 10px;
}
div[data-testid="stMetric"] {
    background-color: #ffffff;
    border-left: 5px solid #2B6CB0;
    padding: 15px;
    border-radius: 8px;
    box-shadow: 0 2px 4px rgba(0,0,0,0.05);
}
.stTabs [data-baseweb="tab-list"] { gap: 8px; }
.stTabs [data-baseweb="tab"] {
    height: 45px;
    white-space: pre-wrap;
    background-color: #EDF2F7;
    border-radius: 6px 6px 0 0;
    color: #2D3748;
    font-weight: 600;
    padding: 10px 16px;
}
.stTabs [aria-selected="true"] { background-color: #1A365D !important; color: white !important; }
.stButton>button { border-radius: 6px; font-weight: 600; }
</style>
"""

CABECALHO_HTML = """
<div class="header-container">
    <h1>🚗 DETRAN/MA — Sistema de Gestão de Exames Práticos</h1>
    <p>Planejamento de Vagas, Controle de Efetivo de Examinadores e Distribuição por Localidades</p>
</div>
"""

LOG = logging.getLogger("detran")
Viagem = dict[str, Any]


def _kwargs_largura_total() -> dict[str, Any]:
    """Compatibilidade entre versões do Streamlit.

    ``use_container_width`` foi substituído por ``width="stretch"``. Detectamos
    a assinatura em tempo de execução para o arquivo funcionar tanto em
    instalações antigas quanto nas novas, sem emitir avisos de depreciação.
    """
    try:
        parametros = inspect.signature(st.dataframe).parameters
    except (TypeError, ValueError):
        return {"use_container_width": True}
    if "width" in parametros:
        return {"width": "stretch"}
    return {"use_container_width": True}


LARGURA_TOTAL = _kwargs_largura_total()


# ===========================================================================
# 2. REGRAS DE NEGÓCIO
# ===========================================================================
# Nenhuma função desta seção acessa session_state ou banco: todas recebem os
# dados por argumento, o que as torna testáveis isoladamente.

_PADRAO_ISO = re.compile(r"^\s*(\d{4})-(\d{1,2})-(\d{1,2})")
_PADRAO_BR = re.compile(r"^\s*(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?")


def para_data(valor: Any) -> datetime.date | None:
    """Normaliza qualquer representação de data para ``datetime.date``."""
    if valor is None or valor == "":
        return None
    if isinstance(valor, datetime.datetime):
        return valor.date()
    if isinstance(valor, datetime.date):
        return valor
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.datetime.strptime(texto[:10], formato).date()
        except ValueError:
            continue
    LOG.warning("Não foi possível interpretar a data %r", valor)
    return None


def formatar_br(data: datetime.date | None) -> str:
    return data.strftime("%d/%m/%Y") if data else "-"


def formatar_curto(data: datetime.date | None) -> str:
    return data.strftime("%d/%m") if data else "-"


def semana_chave(data: datetime.date) -> str:
    segunda = data - datetime.timedelta(days=data.weekday())
    return segunda.isoformat()


def nome_dia_semana(data: datetime.date) -> str:
    return DIAS_SEMANA_OPCOES[data.weekday()]


def dias_do_intervalo(
    inicio: datetime.date, fim: datetime.date, apenas_uteis: bool = False
) -> list[datetime.date]:
    dias = []
    cursor = inicio
    while cursor <= fim:
        if not apenas_uteis or cursor.weekday() < 5:
            dias.append(cursor)
        cursor += datetime.timedelta(days=1)
    return dias


def _data_ou_none(ano: int, mes: int, dia: int) -> datetime.date | None:
    try:
        return datetime.date(ano, mes, dia)
    except ValueError:
        return None


@lru_cache(maxsize=64)
def feriados_estaduais(ano: int, uf: str = UF_FERIADOS) -> frozenset[datetime.date]:
    """Feriados nacionais e estaduais, tolerante à versão da biblioteca."""
    try:
        calendario = holidays.Brazil(subdiv=uf, years=ano)
    except TypeError:
        try:
            calendario = holidays.BR(state=uf, years=ano)
        except Exception:
            LOG.exception("Falha ao carregar feriados de %s/%s", uf, ano)
            return frozenset()
    except Exception:
        LOG.exception("Falha ao carregar feriados de %s/%s", uf, ano)
        return frozenset()
    return frozenset(calendario.keys())


def extrair_data_feriado(linha: str, ano_referencia: int) -> datetime.date | None:
    """Lê ``25/12 - Natal``, ``25/12/2027`` ou ``2026-06-15`` e devolve a data."""
    if not linha or not linha.strip():
        return None
    texto = linha.strip()

    iso = _PADRAO_ISO.match(texto)
    if iso:
        return _data_ou_none(*(int(g) for g in iso.groups()))

    br = _PADRAO_BR.match(texto)
    if not br:
        return None
    dia, mes, ano = br.groups()
    if ano is None:
        ano_int = int(ano_referencia)
    else:
        ano_int = int(ano)
        if ano_int < 100:
            ano_int += 2000
    return _data_ou_none(ano_int, int(mes), int(dia))


def datas_de_feriado_do_texto(texto: str, ano_referencia: int) -> set[datetime.date]:
    datas: set[datetime.date] = set()
    for linha in str(texto or "").splitlines():
        data = extrair_data_feriado(linha.strip(), ano_referencia)
        if data:
            datas.add(data)
    return datas


# --- Chaves compostas -------------------------------------------------------
def _sanitizar(texto: str) -> str:
    return str(texto).replace(SEPARADOR_CHAVE, "/")


def chave_historico(banca: str, mes: str, ano: int | str, local: str) -> str:
    return SEPARADOR_CHAVE.join(
        [_sanitizar(banca), _sanitizar(mes), str(ano), _sanitizar(local)]
    )


def desmontar_chave_historico(chave: str) -> tuple[str, str, str, str] | None:
    partes = chave.split(SEPARADOR_CHAVE)
    if len(partes) != 4:
        return None
    return partes[0], partes[1], partes[2], partes[3]


def chave_local(banca: str, local: str) -> str:
    return SEPARADOR_CHAVE.join([_sanitizar(banca), _sanitizar(local)])


def chave_disponibilidade(banca: str, data: datetime.date) -> str:
    return SEPARADOR_CHAVE.join([_sanitizar(banca), semana_chave(data)])


def normalizar_texto(texto: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", str(texto))
    sem_acento = "".join(c for c in sem_acento if not unicodedata.combining(c))
    return sem_acento.strip().lower()


# --- Estrutura organizacional ----------------------------------------------
def obter_sede(bancas_config: dict[str, list[str]], banca: str) -> str | None:
    localidades = bancas_config.get(banca) or []
    return localidades[0] if localidades else None


def todas_localidades(bancas_config: dict[str, list[str]]) -> list[str]:
    encontradas: set[str] = set()
    for lista in bancas_config.values():
        encontradas.update(lista)
    return sorted(encontradas)


# --- Viagens ----------------------------------------------------------------
def bancas_da_viagem(viagem: Viagem) -> list[str]:
    bancas = []
    titular = viagem.get("Banca")
    if titular:
        bancas.append(titular)
    for banca in viagem.get("Bancas Apoio") or []:
        if banca not in bancas:
            bancas.append(banca)
    return bancas


def banca_vinculada(viagem: Viagem, banca: str) -> bool:
    return banca in bancas_da_viagem(viagem)


def examinadores_da_banca_na_viagem(viagem: Viagem, banca: str) -> int:
    if not banca_vinculada(viagem, banca):
        return 0
    por_banca = viagem.get("Examinadores Por Banca") or {}
    if banca in por_banca:
        try:
            return int(por_banca[banca])
        except (TypeError, ValueError):
            return 0
    if viagem.get("Banca") == banca:
        try:
            return int(viagem.get("Examinadores", 0))
        except (TypeError, ValueError):
            return 0
    return 0


def total_examinadores(viagem: Viagem) -> int:
    por_banca = viagem.get("Examinadores Por Banca") or {}
    if por_banca:
        return sum(int(v or 0) for v in por_banca.values())
    try:
        return int(viagem.get("Examinadores", 0))
    except (TypeError, ValueError):
        return 0


def turno_da_viagem_no_dia(viagem: Viagem, data: datetime.date) -> str:
    """Turno efetivo da viagem em um dia específico.

    O turno geral da viagem vale para todo o período, mas cada dia pode ter
    um ajuste próprio. É o que permite uma mesma equipe atender Pinheiro pela
    manhã e São Bento à tarde no dia 12, sem que o sistema conte duas equipes.
    """
    ajustes = viagem.get("Turnos Por Data") or {}
    especifico = ajustes.get(data.isoformat())
    if especifico:
        return especifico
    return viagem.get("Turno", TURNO_INTEGRAL)


def viagem_cobre_periodo(viagem: Viagem, data: datetime.date, periodo: str) -> bool:
    inicio = para_data(viagem.get("Data Inicio"))
    fim = para_data(viagem.get("Data Fim"))
    if not inicio or not fim or not (inicio <= data <= fim):
        return False
    turno = turno_da_viagem_no_dia(viagem, data)
    if turno == TURNO_SEM_ATENDIMENTO:
        return False
    return turno == TURNO_INTEGRAL or turno == periodo


def viagem_cobre_dia(viagem: Viagem, data: datetime.date) -> bool:
    return viagem_cobre_periodo(viagem, data, TURNO_MANHA) or viagem_cobre_periodo(
        viagem, data, TURNO_TARDE
    )


def dias_efetivos_da_viagem(viagem: Viagem) -> list[datetime.date]:
    """Dias em que a equipe realmente atende, já descontando 'Não atende'."""
    inicio = para_data(viagem.get("Data Inicio"))
    fim = para_data(viagem.get("Data Fim"))
    if not inicio or not fim:
        return []
    return [d for d in dias_do_intervalo(inicio, fim) if viagem_cobre_dia(viagem, d)]


def descricao_turnos(viagem: Viagem) -> str:
    """Texto curto com os dias que fogem do turno geral da viagem."""
    ajustes = viagem.get("Turnos Por Data") or {}
    if not ajustes:
        return ""
    partes = []
    for iso in sorted(ajustes):
        data = para_data(iso)
        turno = ajustes[iso]
        if not data or turno == viagem.get("Turno", TURNO_INTEGRAL):
            continue
        partes.append(f"{formatar_curto(data)}: {turno.lower()}")
    return " · ".join(partes)


def periodo_do_horario(horario: str) -> str:
    try:
        hora, minuto = (int(p) for p in str(horario).split(":")[:2])
    except (TypeError, ValueError):
        LOG.warning("Horário inválido: %r — assumindo Manhã", horario)
        return TURNO_MANHA
    return TURNO_MANHA if (hora, minuto) < (12, 0) else TURNO_TARDE


# --- Ajuste 3: sugestão de horário a partir do deslocamento -----------------
def _hhmm_para_minutos(horario: str) -> int | None:
    try:
        hora, minuto = (int(p) for p in str(horario).split(":")[:2])
        return hora * 60 + minuto
    except (TypeError, ValueError):
        return None


def _minutos_para_horario(minutos: int) -> datetime.time:
    minutos = max(0, min(int(minutos), 23 * 60 + 59))
    return datetime.time(minutos // 60, minutos % 60)


def tempo_deslocamento_padrao(banca: str, local: str) -> int:
    """Minutos estimados de deslocamento da sede da banca até o local.

    0 quando a localidade não precisa de deslocamento (sede, região
    metropolitana) ou quando não há estimativa cadastrada para ela. Se o
    local não está listado sob esta banca (situação de banca atendendo um
    município que normalmente é de outra), reaproveita a estimativa de onde
    ele costuma ser atendido — mas o valor final sempre pode ser corrigido
    pelo usuário no cadastro da viagem.
    """
    direto = TEMPO_DESLOCAMENTO_PADRAO_MIN.get(banca, {}).get(local)
    if direto is not None:
        return direto
    for mapa in TEMPO_DESLOCAMENTO_PADRAO_MIN.values():
        if local in mapa:
            return mapa[local]
    return 0


def horario_chegada_estimado(
    minutos_deslocamento: int, partida: datetime.time = HORA_PARTIDA_PADRAO
) -> datetime.time:
    """Horário de chegada saindo da sede às `partida` mais o deslocamento."""
    total = partida.hour * 60 + partida.minute + max(0, int(minutos_deslocamento))
    return _minutos_para_horario(total)


def sugerir_horario_inicio(
    horarios_ativos: Sequence[str], chegada: datetime.time
) -> str | None:
    """Primeiro horário da grade que já acomoda a chegada estimada."""
    chegada_min = chegada.hour * 60 + chegada.minute
    candidatos = sorted(
        (h for h in horarios_ativos if _hhmm_para_minutos(h) is not None),
        key=lambda h: _hhmm_para_minutos(h),
    )
    for horario in candidatos:
        if _hhmm_para_minutos(horario) >= chegada_min:
            return horario
    return None


def horario_fim_turno(
    horarios_ativos: Sequence[str], periodo: str
) -> datetime.time | None:
    """Horário estimado de término do turno, assumindo turmas de 1h."""
    do_periodo = [h for h in horarios_ativos if periodo_do_horario(h) == periodo]
    if not do_periodo:
        return None
    ultimo = max(do_periodo, key=lambda h: _hhmm_para_minutos(h))
    return _minutos_para_horario(_hhmm_para_minutos(ultimo) + DURACAO_TURMA_MIN)


def sugerir_horario_segunda_cidade(
    horarios_ativos: Sequence[str],
    fim_turno_cidade1: datetime.time,
    deslocamento_entre_cidades_min: int,
) -> tuple[str | None, datetime.time]:
    """Sugestão de início na 2ª cidade da mesma viagem, no mesmo dia.

    Conta o fim do turno da primeira cidade, mais 1h de almoço obrigatória,
    mais o deslocamento entre as duas cidades. Devolve (horário sugerido na
    grade, horário de chegada estimado) — o horário sugerido pode vir None se
    a chegada estimada for depois do último horário cadastrado.
    """
    total = (
        fim_turno_cidade1.hour * 60
        + fim_turno_cidade1.minute
        + DURACAO_ALMOCO_MIN
        + max(0, int(deslocamento_entre_cidades_min))
    )
    chegada = _minutos_para_horario(total)
    return sugerir_horario_inicio(horarios_ativos, chegada), chegada


def viagens_ativas_no_periodo(
    viagens: Sequence[Viagem],
    banca: str,
    data: datetime.date,
    periodo: str,
    ignorar_id: Any = None,
) -> list[Viagem]:
    ativas = []
    for v in viagens:
        if ignorar_id is not None and v.get("id") == ignorar_id:
            continue
        if banca_vinculada(v, banca) and viagem_cobre_periodo(v, data, periodo):
            ativas.append(v)
    return ativas


def viagem_do_local(
    viagens: Sequence[Viagem],
    banca: str,
    local: str,
    data: datetime.date,
    periodo: str,
) -> Viagem | None:
    for v in viagens:
        if (
            banca_vinculada(v, banca)
            and v.get("Destino") == local
            and viagem_cobre_periodo(v, data, periodo)
        ):
            return v
    return None


def viagem_do_local_no_dia(
    viagens: Sequence[Viagem], banca: str, local: str, data: datetime.date
) -> Viagem | None:
    for v in viagens:
        if (
            banca_vinculada(v, banca)
            and v.get("Destino") == local
            and viagem_cobre_dia(v, data)
        ):
            return v
    return None


def conflitos_de_rota(
    viagens: Sequence[Viagem],
    banca: str,
    numero_equipe: int,
    inicio: datetime.date,
    fim: datetime.date,
    turno: str,
    turnos_por_data: dict[str, str] | None = None,
    ignorar_id: Any = None,
) -> list[Viagem]:
    """Trechos da mesma equipe que disputam o mesmo turno no mesmo dia.

    A checagem passou a ser dia a dia: antes bastava haver sobreposição de
    intervalo para bloquear, o que impedia justamente o caso de Pinheiro pela
    manhã e São Bento à tarde no mesmo dia.
    """
    candidata = {
        "Data Inicio": inicio,
        "Data Fim": fim,
        "Turno": turno,
        "Turnos Por Data": turnos_por_data or {},
    }
    conflitos: list[Viagem] = []

    for v in viagens:
        if ignorar_id is not None and v.get("id") == ignorar_id:
            continue
        if int(v.get("Numero Banca Itinerante", 0) or 0) != int(numero_equipe):
            continue
        if not banca_vinculada(v, banca):
            continue
        v_ini = para_data(v.get("Data Inicio"))
        v_fim = para_data(v.get("Data Fim"))
        if not v_ini or not v_fim:
            continue

        for dia in dias_do_intervalo(max(inicio, v_ini), min(fim, v_fim)):
            for periodo in (TURNO_MANHA, TURNO_TARDE):
                if viagem_cobre_periodo(candidata, dia, periodo) and (
                    viagem_cobre_periodo(v, dia, periodo)
                ):
                    conflitos.append(v)
                    break
            else:
                continue
            break
    return conflitos


def numeros_de_equipe(viagens: Sequence[Viagem], banca: str) -> list[int]:
    return sorted(
        {
            int(v.get("Numero Banca Itinerante", 0) or 0)
            for v in viagens
            if v.get("Banca") == banca
        }
        - {0}
    )


def proximo_numero_equipe(viagens: Sequence[Viagem], banca: str) -> int:
    numeros = numeros_de_equipe(viagens, banca)
    return max(numeros) + 1 if numeros else 1


def garantir_numeracao_equipes(viagens: list[Viagem]) -> list[Viagem]:
    contador: dict[str, int] = {}
    for v in viagens:
        banca = v.get("Banca", "")
        numero = v.get("Numero Banca Itinerante")
        if not numero:
            contador[banca] = contador.get(banca, 0) + 1
            v["Numero Banca Itinerante"] = contador[banca]
        else:
            contador[banca] = max(contador.get(banca, 0), int(numero))
    return viagens


# --- Efetivo ----------------------------------------------------------------
def equipes_em_viagem(
    viagens: Sequence[Viagem], banca: str, data: datetime.date
) -> dict[str, dict[int, int]]:
    """Examinadores em viagem por período, agrupados por equipe itinerante.

    Uma equipe que atende dois municípios no mesmo dia em turnos diferentes é
    contada uma única vez — daí o ``max`` por número de equipe.
    """
    resultado: dict[str, dict[int, int]] = {TURNO_MANHA: {}, TURNO_TARDE: {}}
    for v in viagens:
        if not banca_vinculada(v, banca):
            continue
        quantidade = examinadores_da_banca_na_viagem(v, banca)
        if quantidade <= 0:
            continue
        equipe = int(v.get("Numero Banca Itinerante", 0) or 0)
        for periodo in (TURNO_MANHA, TURNO_TARDE):
            if viagem_cobre_periodo(v, data, periodo):
                atual = resultado[periodo].get(equipe, 0)
                resultado[periodo][equipe] = max(atual, quantidade)
    return resultado


def efetivo_em_viagem(
    viagens: Sequence[Viagem], banca: str, data: datetime.date
) -> tuple[int, int]:
    equipes = equipes_em_viagem(viagens, banca, data)
    return sum(equipes[TURNO_MANHA].values()), sum(equipes[TURNO_TARDE].values())


def pico_diario(manha: int, tarde: int) -> int:
    """O efetivo necessário no dia é o maior dos turnos, nunca a soma."""
    return max(manha, tarde)


def agrupar_trechos_por_equipe(
    viagens: Sequence[Viagem], banca: str | None = None
) -> list[dict[str, Any]]:
    """Consolida os trechos de cada equipe itinerante em um único registro."""
    grupos: dict[tuple[str, int], list[Viagem]] = {}
    for v in viagens:
        banca_v = v.get("Banca", "")
        if banca and banca_v != banca:
            continue
        chave = (banca_v, int(v.get("Numero Banca Itinerante", 1) or 1))
        grupos.setdefault(chave, []).append(v)

    consolidados = []
    for (banca_v, numero), trechos in sorted(grupos.items()):
        trechos = sorted(trechos, key=lambda t: para_data(t["Data Inicio"]))
        inicio = min(para_data(t["Data Inicio"]) for t in trechos)
        fim = max(para_data(t["Data Fim"]) for t in trechos)
        destinos: list[str] = []
        for t in trechos:
            destino = t.get("Destino", "")
            if destino and destino not in destinos:
                destinos.append(destino)
        consolidados.append(
            {
                "Banca": banca_v,
                "Numero": numero,
                "Trechos": trechos,
                "Inicio": inicio,
                "Fim": fim,
                "Destinos": destinos,
                "Examinadores": max(
                    (total_examinadores(t) for t in trechos), default=0
                ),
            }
        )
    return consolidados


def equipes_ativas_no_intervalo(
    viagens: Sequence[Viagem],
    banca: str,
    inicio: datetime.date,
    fim: datetime.date,
) -> list[dict[str, Any]]:
    """Equipes itinerantes da banca que atuam dentro do intervalo informado."""
    ativas = []
    for grupo in agrupar_trechos_por_equipe(viagens):
        if not any(banca_vinculada(t, banca) for t in grupo["Trechos"]):
            continue
        dias = [
            d
            for t in grupo["Trechos"]
            for d in dias_efetivos_da_viagem(t)
            if inicio <= d <= fim
        ]
        if not dias:
            continue
        copia = dict(grupo)
        copia["Inicio Janela"] = min(dias)
        copia["Fim Janela"] = max(dias)
        copia["Examinadores"] = max(
            (
                examinadores_da_banca_na_viagem(t, banca)
                for t in grupo["Trechos"]
            ),
            default=0,
        )
        ativas.append(copia)
    return ativas


def formatar_datas_exames(dias_do_mes: Iterable[int]) -> str:
    """Agrupa [3,4,5,10] em '03 a 05; 10' para a publicação oficial."""
    dias = sorted({int(d) for d in dias_do_mes})
    if not dias:
        return "-"
    grupos: list[list[int]] = []
    atual = [dias[0]]
    for dia in dias[1:]:
        if dia == atual[-1] + 1:
            atual.append(dia)
        else:
            grupos.append(atual)
            atual = [dia]
    grupos.append(atual)

    partes = []
    for grupo in grupos:
        if len(grupo) == 1:
            partes.append(f"{grupo[0]:02d}")
        elif len(grupo) == 2:
            partes.append(f"{grupo[0]:02d} e {grupo[1]:02d}")
        else:
            partes.append(f"{grupo[0]:02d} a {grupo[-1]:02d}")
    return "; ".join(partes)


def bancas_para_validar(banca: str, bancas_apoio: Sequence[str] | None) -> list[str]:
    bancas = [banca]
    for outra in bancas_apoio or []:
        if outra not in bancas:
            bancas.append(outra)
    return bancas


def bancas_apoio_disponiveis(banca: str) -> list[str]:
    if banca not in BANCAS_APOIO_CONJUNTO:
        return []
    return [b for b in BANCAS_APOIO_CONJUNTO if b != banca]


# --- Ordenação para os relatórios em PDF ------------------------------------
def _primeira_data_lancada(df: pd.DataFrame) -> datetime.date:
    """Data do primeiro exame disponível da localidade (para ordenar)."""
    if df is None or df.empty:
        return datetime.date.max
    disponiveis = df[(df["Status"] == STATUS_DISPONIVEL) & (df["Total"] > 0)]
    if disponiveis.empty:
        return datetime.date.max
    datas = [para_data(d) for d in disponiveis["Data"].unique()]
    validas = [d for d in datas if d]
    return min(validas) if validas else datetime.date.max


def ordenar_localidades_para_pdf(
    banca: str,
    bancas_config: dict[str, list[str]],
    dados_por_local: dict[str, pd.DataFrame],
) -> list[tuple[str, str]]:
    """Ordem de apresentação das localidades no PDF da banca.

    Regra: unidades da sede primeiro, depois a região metropolitana e, por
    fim, os demais municípios pela data do primeiro exame lançado.
    Devolve pares (grupo, localidade).
    """
    prioridade = ORDEM_PRIORITARIA_PDF.get(banca)
    if not prioridade:
        sede = obter_sede(bancas_config, banca)
        prioridade = [sede] if sede else []

    indice_prioridade = {normalizar_texto(nome): i for i, nome in enumerate(prioridade)}

    prioritarias: list[tuple[int, str]] = []
    demais: list[tuple[datetime.date, str, str]] = []

    for local, df in dados_por_local.items():
        chave = normalizar_texto(local)
        if chave in indice_prioridade:
            prioritarias.append((indice_prioridade[chave], local))
        else:
            demais.append((_primeira_data_lancada(df), normalizar_texto(local), local))

    ordenadas: list[tuple[str, str]] = []
    for _, local in sorted(prioritarias):
        grupo = GRUPOS_PDF.get(local)
        if grupo is None:
            grupo = GRUPO_SEDE_GENERICO
        ordenadas.append((grupo, local))
    for _, _, local in sorted(demais):
        ordenadas.append((GRUPO_INTERIOR, local))
    return ordenadas


# ===========================================================================
# 3. PERSISTÊNCIA
# ===========================================================================
# SQLite com gravação por entidade e só quando o conteúdo muda de fato.

MAX_SNAPSHOTS = 20
_LOCK = threading.Lock()


class BackupInvalido(Exception):
    """Arquivo de backup fora do formato esperado."""


def caminho_banco() -> Path:
    return Path(os.environ.get("DETRAN_DB", "dados/detran.db"))


def pasta_snapshots() -> Path:
    return Path(os.environ.get("DETRAN_SNAPSHOTS", "dados/snapshots"))


ESQUEMA = """
CREATE TABLE IF NOT EXISTS config (
    chave TEXT PRIMARY KEY,
    valor TEXT NOT NULL,
    atualizado_em TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS bancas (
    nome  TEXT PRIMARY KEY,
    ordem INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE IF NOT EXISTS localidades (
    banca TEXT NOT NULL,
    nome  TEXT NOT NULL,
    ordem INTEGER NOT NULL DEFAULT 0,
    PRIMARY KEY (banca, nome)
);

CREATE TABLE IF NOT EXISTS viagens (
    id                     INTEGER PRIMARY KEY AUTOINCREMENT,
    banca                  TEXT NOT NULL,
    numero_equipe          INTEGER NOT NULL DEFAULT 1,
    destino                TEXT NOT NULL,
    data_inicio            TEXT NOT NULL,
    data_fim               TEXT NOT NULL,
    turno                  TEXT NOT NULL DEFAULT 'Dia inteiro',
    turnos_por_data        TEXT NOT NULL DEFAULT '{}',
    examinadores           INTEGER NOT NULL DEFAULT 0,
    bancas_apoio           TEXT NOT NULL DEFAULT '[]',
    examinadores_por_banca TEXT NOT NULL DEFAULT '{}',
    observacoes            TEXT NOT NULL DEFAULT ''
);

CREATE TABLE IF NOT EXISTS historico (
    chave      TEXT NOT NULL,
    data       TEXT NOT NULL,
    horario    TEXT NOT NULL,
    dia_semana TEXT NOT NULL DEFAULT '',
    status     TEXT NOT NULL,
    exam_m     INTEGER NOT NULL DEFAULT 0,
    exam_t     INTEGER NOT NULL DEFAULT 0,
    vagas      TEXT NOT NULL DEFAULT '{}',
    total      INTEGER NOT NULL DEFAULT 0,
    PRIMARY KEY (chave, data, horario)
);

CREATE TABLE IF NOT EXISTS dias_permitidos (
    chave TEXT PRIMARY KEY,
    dias  TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS feriados_locais (
    chave TEXT PRIMARY KEY,
    texto TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS disponibilidade_semanal (
    chave TEXT PRIMARY KEY,
    valor INTEGER NOT NULL
);

CREATE TABLE IF NOT EXISTS deslocamentos (
    chave   TEXT PRIMARY KEY,
    minutos INTEGER NOT NULL
);
"""


def _migrar_esquema(conexao: sqlite3.Connection) -> None:
    """Adiciona colunas novas em bancos criados por versões anteriores."""
    colunas = {
        linha["name"] for linha in conexao.execute("PRAGMA table_info(viagens)")
    }
    if colunas and "turnos_por_data" not in colunas:
        LOG.info("Migrando esquema: adicionando viagens.turnos_por_data")
        with conexao:
            conexao.execute(
                "ALTER TABLE viagens ADD COLUMN turnos_por_data TEXT NOT NULL"
                " DEFAULT '{}'"
            )


@st.cache_resource(show_spinner=False)
def _abrir_conexao(caminho: str) -> sqlite3.Connection:
    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    conexao = sqlite3.connect(caminho, check_same_thread=False)
    conexao.row_factory = sqlite3.Row
    conexao.execute("PRAGMA journal_mode=WAL")
    conexao.execute("PRAGMA synchronous=NORMAL")
    with conexao:
        conexao.executescript(ESQUEMA)
    _migrar_esquema(conexao)
    LOG.info("Banco aberto em %s", Path(caminho).resolve())
    return conexao


def conectar() -> sqlite3.Connection:
    return _abrir_conexao(str(caminho_banco()))


# --- Utilitários ------------------------------------------------------------
def _json_seguro(valor: Any) -> Any:
    if isinstance(valor, (datetime.date, datetime.datetime)):
        return valor.isoformat()
    if hasattr(valor, "item"):
        return valor.item()
    if isinstance(valor, (set, frozenset)):
        return sorted(valor)
    return str(valor)


def _dump(valor: Any) -> str:
    return json.dumps(valor, ensure_ascii=False, default=_json_seguro, allow_nan=False)


def _inteiro(valor: Any, padrao: int = 0) -> int:
    try:
        if valor is None or (isinstance(valor, float) and pd.isna(valor)):
            return padrao
        return int(valor)
    except (TypeError, ValueError):
        return padrao


def _assinatura(valor: Any) -> str:
    try:
        bruto = _dump(valor)
    except (TypeError, ValueError):
        bruto = repr(valor)
    return hashlib.sha1(bruto.encode("utf-8")).hexdigest()


def _mudou(nome: str, valor: Any) -> bool:
    """Evita reescrever o banco a cada clique de widget."""
    assinaturas = st.session_state.setdefault("_assinaturas", {})
    nova = _assinatura(valor)
    if assinaturas.get(nome) == nova:
        return False
    assinaturas[nome] = nova
    return True


def _registrar_erro(mensagem: str, excecao: Exception) -> None:
    LOG.exception(mensagem)
    st.session_state["erro_persistencia"] = f"{mensagem}: {excecao}"


def _marcar_sucesso() -> None:
    st.session_state["erro_persistencia"] = None
    st.session_state["ultimo_salvamento"] = datetime.datetime.now().strftime(
        "%d/%m/%Y %H:%M:%S"
    )


# --- Leitura ----------------------------------------------------------------
def _ler_config(conexao: sqlite3.Connection, chave: str, padrao: Any) -> Any:
    linha = conexao.execute(
        "SELECT valor FROM config WHERE chave = ?", (chave,)
    ).fetchone()
    if not linha:
        return padrao
    try:
        return json.loads(linha["valor"])
    except json.JSONDecodeError:
        LOG.warning("Config %s corrompida — usando padrão", chave)
        return padrao


def _ler_bancas(conexao: sqlite3.Connection) -> dict[str, list[str]]:
    bancas: dict[str, list[str]] = {}
    for linha in conexao.execute("SELECT nome FROM bancas ORDER BY ordem, nome"):
        bancas[linha["nome"]] = []
    for linha in conexao.execute(
        "SELECT banca, nome FROM localidades ORDER BY banca, ordem, nome"
    ):
        bancas.setdefault(linha["banca"], []).append(linha["nome"])
    return bancas


def _ler_viagens(conexao: sqlite3.Connection) -> list[Viagem]:
    viagens = []
    for linha in conexao.execute("SELECT * FROM viagens ORDER BY data_inicio, id"):
        try:
            bancas_apoio = json.loads(linha["bancas_apoio"])
            por_banca = json.loads(linha["examinadores_por_banca"])
        except json.JSONDecodeError:
            bancas_apoio, por_banca = [], {}
        try:
            turnos = json.loads(linha["turnos_por_data"])
        except (json.JSONDecodeError, IndexError, KeyError):
            turnos = {}
        viagens.append(
            {
                "id": linha["id"],
                "Banca": linha["banca"],
                "Numero Banca Itinerante": linha["numero_equipe"],
                "Destino": linha["destino"],
                "Data Inicio": para_data(linha["data_inicio"]),
                "Data Fim": para_data(linha["data_fim"]),
                "Turno": linha["turno"],
                "Turnos Por Data": turnos,
                "Examinadores": linha["examinadores"],
                "Bancas Apoio": bancas_apoio,
                "Examinadores Por Banca": por_banca,
                "Observações": linha["observacoes"],
            }
        )
    return viagens


def _ler_historico(conexao: sqlite3.Connection) -> dict[str, pd.DataFrame]:
    por_chave: dict[str, list[dict[str, Any]]] = {}
    for linha in conexao.execute(
        "SELECT * FROM historico ORDER BY chave, data, horario"
    ):
        try:
            vagas = json.loads(linha["vagas"])
        except json.JSONDecodeError:
            vagas = {}
        registro = {
            "Data": linha["data"],
            "Dia da Semana": linha["dia_semana"],
            "Status": linha["status"],
            "Exam. M": linha["exam_m"],
            "Exam. T": linha["exam_t"],
            "Horário": linha["horario"],
        }
        for coluna in COLS_VAGAS:
            registro[coluna] = _inteiro(vagas.get(coluna, 0))
        registro["Total"] = linha["total"]
        por_chave.setdefault(linha["chave"], []).append(registro)
    return {chave: pd.DataFrame(linhas) for chave, linhas in por_chave.items()}


def _ler_mapa(conexao: sqlite3.Connection, tabela: str, coluna: str) -> dict[str, Any]:
    resultado = {}
    for linha in conexao.execute(f"SELECT chave, {coluna} FROM {tabela}"):
        valor = linha[coluna]
        if coluna == "dias":
            try:
                valor = json.loads(valor)
            except json.JSONDecodeError:
                valor = []
        resultado[linha["chave"]] = valor
    return resultado


# --- Escrita ----------------------------------------------------------------
def salvar_config(chave: str, valor: Any, forcar: bool = False) -> None:
    if not forcar and not _mudou(f"config:{chave}", valor):
        return
    conexao = conectar()
    try:
        with _LOCK, conexao:
            conexao.execute(
                "INSERT INTO config (chave, valor, atualizado_em) VALUES (?, ?, ?) "
                "ON CONFLICT(chave) DO UPDATE SET valor = excluded.valor, "
                "atualizado_em = excluded.atualizado_em",
                (chave, _dump(valor), datetime.datetime.now().isoformat()),
            )
        _marcar_sucesso()
    except Exception as erro:
        _registrar_erro(f"Falha ao salvar a configuração '{chave}'", erro)


def salvar_bancas(bancas_config: dict[str, list[str]], forcar: bool = False) -> None:
    if not forcar and not _mudou("bancas", bancas_config):
        return
    conexao = conectar()
    try:
        with _LOCK, conexao:
            conexao.execute("DELETE FROM bancas")
            conexao.execute("DELETE FROM localidades")
            for ordem_banca, (banca, locais) in enumerate(bancas_config.items()):
                conexao.execute(
                    "INSERT INTO bancas (nome, ordem) VALUES (?, ?)",
                    (banca, ordem_banca),
                )
                conexao.executemany(
                    "INSERT INTO localidades (banca, nome, ordem) VALUES (?, ?, ?)",
                    [(banca, local, i) for i, local in enumerate(locais)],
                )
        _marcar_sucesso()
    except Exception as erro:
        _registrar_erro("Falha ao salvar bancas e localidades", erro)


def salvar_viagens(viagens: list[Viagem], forcar: bool = False) -> None:
    comparavel = [{k: v for k, v in viagem.items() if k != "id"} for viagem in viagens]
    if not forcar and not _mudou("viagens", comparavel):
        return
    conexao = conectar()
    try:
        with _LOCK, conexao:
            conexao.execute("DELETE FROM viagens")
            for viagem in viagens:
                cursor = conexao.execute(
                    "INSERT INTO viagens (banca, numero_equipe, destino, data_inicio,"
                    " data_fim, turno, turnos_por_data, examinadores, bancas_apoio,"
                    " examinadores_por_banca, observacoes)"
                    " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (
                        viagem.get("Banca", ""),
                        _inteiro(viagem.get("Numero Banca Itinerante"), 1),
                        viagem.get("Destino", ""),
                        para_data(viagem.get("Data Inicio")).isoformat(),
                        para_data(viagem.get("Data Fim")).isoformat(),
                        viagem.get("Turno", TURNO_INTEGRAL),
                        _dump(viagem.get("Turnos Por Data") or {}),
                        _inteiro(viagem.get("Examinadores")),
                        _dump(viagem.get("Bancas Apoio") or []),
                        _dump(viagem.get("Examinadores Por Banca") or {}),
                        str(viagem.get("Observações", "")),
                    ),
                )
                viagem["id"] = cursor.lastrowid
        _marcar_sucesso()
    except Exception as erro:
        _registrar_erro("Falha ao salvar as viagens", erro)


def salvar_historico(chave: str, df: pd.DataFrame, forcar: bool = False) -> None:
    if df is None or df.empty:
        return
    registros = df.to_dict(orient="records")
    if not forcar and not _mudou(f"historico:{chave}", registros):
        return
    conexao = conectar()
    try:
        linhas = []
        for registro in registros:
            vagas = {c: _inteiro(registro.get(c, 0)) for c in COLS_VAGAS}
            linhas.append(
                (
                    chave,
                    str(registro.get("Data", "")),
                    str(registro.get("Horário", "")),
                    str(registro.get("Dia da Semana", "")),
                    str(registro.get("Status", STATUS_DISPONIVEL)),
                    _inteiro(registro.get("Exam. M")),
                    _inteiro(registro.get("Exam. T")),
                    _dump(vagas),
                    _inteiro(registro.get("Total")),
                )
            )
        with _LOCK, conexao:
            conexao.execute("DELETE FROM historico WHERE chave = ?", (chave,))
            conexao.executemany(
                "INSERT INTO historico (chave, data, horario, dia_semana, status,"
                " exam_m, exam_t, vagas, total)"
                " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                linhas,
            )
        _marcar_sucesso()
    except Exception as erro:
        _registrar_erro(f"Falha ao salvar o calendário '{chave}'", erro)


def _salvar_mapa(
    tabela: str, coluna: str, dados: dict[str, Any], nome_cache: str, forcar: bool
) -> None:
    if not forcar and not _mudou(nome_cache, dados):
        return
    conexao = conectar()
    try:
        with _LOCK, conexao:
            conexao.execute(f"DELETE FROM {tabela}")
            conexao.executemany(
                f"INSERT INTO {tabela} (chave, {coluna}) VALUES (?, ?)",
                [
                    (chave, _dump(valor) if coluna == "dias" else valor)
                    for chave, valor in dados.items()
                ],
            )
        _marcar_sucesso()
    except Exception as erro:
        _registrar_erro(f"Falha ao salvar a tabela '{tabela}'", erro)


def salvar_dias_permitidos(dados: dict[str, list[str]], forcar: bool = False) -> None:
    _salvar_mapa("dias_permitidos", "dias", dados, "dias_permitidos", forcar)


def salvar_feriados_locais(dados: dict[str, str], forcar: bool = False) -> None:
    _salvar_mapa("feriados_locais", "texto", dados, "feriados_locais", forcar)


def salvar_disponibilidade(dados: dict[str, int], forcar: bool = False) -> None:
    _salvar_mapa("disponibilidade_semanal", "valor", dados, "disponibilidade", forcar)


def salvar_tempo_deslocamento(dados: dict[str, int], forcar: bool = False) -> None:
    _salvar_mapa("deslocamentos", "minutos", dados, "deslocamentos", forcar)


def remover_historico_do_local(banca: str, local: str) -> None:
    """Limpeza em cascata do histórico de uma localidade."""
    conexao = conectar()
    alvos = [
        chave
        for chave in st.session_state.get("historico_localidades", {})
        if (partes := desmontar_chave_historico(chave))
        and partes[0] == banca
        and partes[3] == local
    ]
    try:
        with _LOCK, conexao:
            for chave in alvos:
                conexao.execute("DELETE FROM historico WHERE chave = ?", (chave,))
        for chave in alvos:
            st.session_state["historico_localidades"].pop(chave, None)
            st.session_state.get("_assinaturas", {}).pop(f"historico:{chave}", None)
        _marcar_sucesso()
    except Exception as erro:
        _registrar_erro(f"Falha ao remover o histórico de '{local}'", erro)


# --- Carga inicial ----------------------------------------------------------
def _semear_padroes(conexao: sqlite3.Connection) -> None:
    existe = conexao.execute("SELECT COUNT(*) AS n FROM bancas").fetchone()["n"]
    if existe:
        return
    LOG.info("Banco vazio — semeando estrutura padrão")
    with _LOCK, conexao:
        for ordem_banca, (banca, locais) in enumerate(BANCAS_PADRAO.items()):
            conexao.execute(
                "INSERT INTO bancas (nome, ordem) VALUES (?, ?)", (banca, ordem_banca)
            )
            conexao.executemany(
                "INSERT INTO localidades (banca, nome, ordem) VALUES (?, ?, ?)",
                [(banca, local, i) for i, local in enumerate(locais)],
            )
        agora = datetime.datetime.now().isoformat()
        for chave, valor in (
            ("lista_horarios", HORARIOS_PADRAO),
            ("horarios_inativos", []),
            ("capacidade_bancas", CAPACIDADE_BANCAS_PADRAO),
            ("limite_fora_sede_bancas", LIMITE_FORA_SEDE_PADRAO),
            ("controle_capacidade_ativo", False),
            ("schema_version", SCHEMA_VERSION),
        ):
            conexao.execute(
                "INSERT OR REPLACE INTO config (chave, valor, atualizado_em)"
                " VALUES (?, ?, ?)",
                (chave, _dump(valor), agora),
            )


def carregar_estado(forcar: bool = False) -> None:
    """Carrega o banco para o session_state uma vez por sessão."""
    if st.session_state.get("_estado_carregado") and not forcar:
        return

    conexao = conectar()
    _semear_padroes(conexao)

    st.session_state["bancas_config"] = _ler_bancas(conexao)
    st.session_state["lista_horarios"] = _ler_config(
        conexao, "lista_horarios", list(HORARIOS_PADRAO)
    )
    st.session_state["horarios_inativos"] = _ler_config(
        conexao, "horarios_inativos", []
    )
    st.session_state["capacidade_bancas"] = {
        **CAPACIDADE_BANCAS_PADRAO,
        **_ler_config(conexao, "capacidade_bancas", {}),
    }
    st.session_state["limite_fora_sede_bancas"] = {
        **LIMITE_FORA_SEDE_PADRAO,
        **_ler_config(conexao, "limite_fora_sede_bancas", {}),
    }
    st.session_state["controle_capacidade_ativo"] = bool(
        _ler_config(conexao, "controle_capacidade_ativo", False)
    )
    st.session_state["viagens_registradas"] = garantir_numeracao_equipes(
        _ler_viagens(conexao)
    )
    st.session_state["historico_localidades"] = _ler_historico(conexao)
    st.session_state["dias_permitidos_dict"] = _ler_mapa(
        conexao, "dias_permitidos", "dias"
    )
    st.session_state["feriados_locais_dict"] = _ler_mapa(
        conexao, "feriados_locais", "texto"
    )
    st.session_state["disponibilidade_semanal"] = {
        chave: _inteiro(valor)
        for chave, valor in _ler_mapa(
            conexao, "disponibilidade_semanal", "valor"
        ).items()
    }
    st.session_state["tempo_deslocamento_dict"] = {
        chave: _inteiro(valor)
        for chave, valor in _ler_mapa(conexao, "deslocamentos", "minutos").items()
    }

    st.session_state["_assinaturas"] = {}
    _mudou("bancas", st.session_state["bancas_config"])
    _mudou(
        "viagens",
        [
            {k: v for k, v in viagem.items() if k != "id"}
            for viagem in st.session_state["viagens_registradas"]
        ],
    )
    _mudou("dias_permitidos", st.session_state["dias_permitidos_dict"])
    _mudou("feriados_locais", st.session_state["feriados_locais_dict"])
    _mudou("disponibilidade", st.session_state["disponibilidade_semanal"])
    _mudou("deslocamentos", st.session_state["tempo_deslocamento_dict"])
    for chave, df in st.session_state["historico_localidades"].items():
        _mudou(f"historico:{chave}", df.to_dict(orient="records"))

    st.session_state["_estado_carregado"] = True
    st.session_state.setdefault("erro_persistencia", None)
    LOG.info(
        "Estado carregado: %d bancas, %d viagens, %d calendários",
        len(st.session_state["bancas_config"]),
        len(st.session_state["viagens_registradas"]),
        len(st.session_state["historico_localidades"]),
    )


def recarregar_do_banco() -> None:
    carregar_estado(forcar=True)


# --- Backup -----------------------------------------------------------------
def montar_backup() -> dict[str, Any]:
    viagens = []
    for viagem in st.session_state.get("viagens_registradas", []):
        copia = dict(viagem)
        copia.pop("id", None)
        copia["Data Inicio"] = para_data(viagem.get("Data Inicio")).isoformat()
        copia["Data Fim"] = para_data(viagem.get("Data Fim")).isoformat()
        viagens.append(copia)

    historico = {
        chave: df.fillna(0).to_dict(orient="records")
        for chave, df in st.session_state.get("historico_localidades", {}).items()
    }

    return {
        "schema_version": SCHEMA_VERSION,
        "gerado_em": datetime.datetime.now().isoformat(),
        "bancas_config": st.session_state.get("bancas_config", {}),
        "lista_horarios": st.session_state.get("lista_horarios", []),
        "horarios_inativos": st.session_state.get("horarios_inativos", []),
        "historico_localidades": historico,
        "feriados_locais_dict": st.session_state.get("feriados_locais_dict", {}),
        "viagens_registradas": viagens,
        "dias_permitidos_dict": st.session_state.get("dias_permitidos_dict", {}),
        "capacidade_bancas": st.session_state.get("capacidade_bancas", {}),
        "limite_fora_sede_bancas": st.session_state.get("limite_fora_sede_bancas", {}),
        "disponibilidade_semanal": st.session_state.get("disponibilidade_semanal", {}),
        "tempo_deslocamento_dict": st.session_state.get("tempo_deslocamento_dict", {}),
        "controle_capacidade_ativo": st.session_state.get(
            "controle_capacidade_ativo", False
        ),
    }


def backup_em_json() -> str:
    return _dump(montar_backup())


CAMPOS_OBRIGATORIOS = {
    "bancas_config": dict,
    "viagens_registradas": list,
    "historico_localidades": dict,
}


def validar_backup(dados: Any) -> dict[str, Any]:
    """Valida a estrutura antes de aplicar. Levanta ``BackupInvalido``."""
    if not isinstance(dados, dict):
        raise BackupInvalido("O arquivo não contém um objeto JSON no nível raiz.")

    versao = dados.get("schema_version")
    if versao is not None and not isinstance(versao, int):
        raise BackupInvalido("O campo 'schema_version' deveria ser um número inteiro.")
    if isinstance(versao, int) and versao > SCHEMA_VERSION:
        raise BackupInvalido(
            f"Backup gerado por uma versão mais nova do sistema (schema {versao};"
            f" esta versão entende até {SCHEMA_VERSION})."
        )

    for campo, tipo in CAMPOS_OBRIGATORIOS.items():
        if campo not in dados:
            raise BackupInvalido(f"Campo obrigatório ausente: '{campo}'.")
        if not isinstance(dados[campo], tipo):
            raise BackupInvalido(
                f"O campo '{campo}' deveria ser {tipo.__name__}, "
                f"veio {type(dados[campo]).__name__}."
            )

    for banca, locais in dados["bancas_config"].items():
        if not isinstance(banca, str) or not isinstance(locais, list):
            raise BackupInvalido(
                "Cada banca em 'bancas_config' precisa mapear para uma lista de"
                " localidades."
            )

    for indice, viagem in enumerate(dados["viagens_registradas"], 1):
        if not isinstance(viagem, dict):
            raise BackupInvalido(f"A viagem nº {indice} não é um objeto.")
        for campo in ("Banca", "Destino", "Data Inicio", "Data Fim"):
            if campo not in viagem:
                raise BackupInvalido(f"A viagem nº {indice} está sem o campo '{campo}'.")
        if para_data(viagem["Data Inicio"]) is None:
            raise BackupInvalido(
                f"A viagem nº {indice} tem data de início ilegível:"
                f" {viagem['Data Inicio']!r}."
            )
        if para_data(viagem["Data Fim"]) is None:
            raise BackupInvalido(
                f"A viagem nº {indice} tem data de término ilegível:"
                f" {viagem['Data Fim']!r}."
            )

    for chave, registros in dados["historico_localidades"].items():
        if not isinstance(registros, list):
            raise BackupInvalido(f"O calendário '{chave}' não é uma lista de linhas.")
        if desmontar_chave_historico(chave) is None:
            raise BackupInvalido(
                f"A chave de calendário '{chave}' está fora do formato esperado"
                " (banca||mês||ano||localidade)."
            )
    return dados


def _migrar_viagem_antiga(viagem: dict[str, Any]) -> Viagem:
    """Converte o formato 'Apoio Conjunto' booleano para a lista explícita."""
    convertida = dict(viagem)
    convertida.pop("id", None)
    convertida["Data Inicio"] = para_data(viagem.get("Data Inicio"))
    convertida["Data Fim"] = para_data(viagem.get("Data Fim"))
    convertida.setdefault("Turnos Por Data", {})

    if "Bancas Apoio" not in convertida:
        convertida["Bancas Apoio"] = []
    if "Examinadores Por Banca" not in convertida:
        por_banca: dict[str, int] = {}
        if viagem.get("Apoio Conjunto"):
            for banca, campo in (
                ("Banca Caxias", "Examinadores Caxias"),
                ("Banca Timon", "Examinadores Timon"),
            ):
                if campo in viagem:
                    por_banca[banca] = _inteiro(viagem[campo])
            convertida["Bancas Apoio"] = [
                b for b in por_banca if b != convertida.get("Banca")
            ]
        convertida["Examinadores Por Banca"] = por_banca
    for campo in ("Apoio Conjunto", "Examinadores Caxias", "Examinadores Timon", "Período"):
        convertida.pop(campo, None)
    return convertida


def aplicar_backup(dados: dict[str, Any]) -> None:
    validar_backup(dados)
    salvar_snapshot(motivo="antes_de_restaurar")

    st.session_state["bancas_config"] = dados["bancas_config"]
    st.session_state["lista_horarios"] = dados.get("lista_horarios") or list(
        HORARIOS_PADRAO
    )
    st.session_state["horarios_inativos"] = dados.get("horarios_inativos", [])
    st.session_state["feriados_locais_dict"] = dados.get("feriados_locais_dict", {})
    st.session_state["dias_permitidos_dict"] = dados.get("dias_permitidos_dict", {})
    st.session_state["capacidade_bancas"] = {
        **CAPACIDADE_BANCAS_PADRAO,
        **dados.get("capacidade_bancas", {}),
    }
    st.session_state["limite_fora_sede_bancas"] = {
        **LIMITE_FORA_SEDE_PADRAO,
        **dados.get("limite_fora_sede_bancas", {}),
    }
    st.session_state["disponibilidade_semanal"] = {
        chave: _inteiro(valor)
        for chave, valor in (dados.get("disponibilidade_semanal") or {}).items()
    }
    st.session_state["tempo_deslocamento_dict"] = {
        chave: _inteiro(valor)
        for chave, valor in (dados.get("tempo_deslocamento_dict") or {}).items()
    }
    st.session_state["controle_capacidade_ativo"] = bool(
        dados.get("controle_capacidade_ativo", False)
    )
    st.session_state["viagens_registradas"] = garantir_numeracao_equipes(
        [_migrar_viagem_antiga(v) for v in dados["viagens_registradas"]]
    )
    st.session_state["historico_localidades"] = {
        chave: pd.DataFrame(registros)
        for chave, registros in dados["historico_localidades"].items()
        if registros
    }

    st.session_state["_assinaturas"] = {}
    salvar_bancas(st.session_state["bancas_config"], forcar=True)
    salvar_viagens(st.session_state["viagens_registradas"], forcar=True)
    salvar_dias_permitidos(st.session_state["dias_permitidos_dict"], forcar=True)
    salvar_feriados_locais(st.session_state["feriados_locais_dict"], forcar=True)
    salvar_disponibilidade(st.session_state["disponibilidade_semanal"], forcar=True)
    salvar_tempo_deslocamento(st.session_state["tempo_deslocamento_dict"], forcar=True)
    for chave, valor in (
        ("lista_horarios", st.session_state["lista_horarios"]),
        ("horarios_inativos", st.session_state["horarios_inativos"]),
        ("capacidade_bancas", st.session_state["capacidade_bancas"]),
        ("limite_fora_sede_bancas", st.session_state["limite_fora_sede_bancas"]),
        ("controle_capacidade_ativo", st.session_state["controle_capacidade_ativo"]),
    ):
        salvar_config(chave, valor, forcar=True)
    for chave, df in st.session_state["historico_localidades"].items():
        salvar_historico(chave, df, forcar=True)


def salvar_snapshot(motivo: str = "manual") -> Path | None:
    """Snapshot rotativo em disco, mantendo os últimos N arquivos."""
    try:
        pasta = pasta_snapshots()
        pasta.mkdir(parents=True, exist_ok=True)
        carimbo = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        destino = pasta / f"snapshot_{carimbo}_{motivo}.json"
        destino.write_text(backup_em_json(), encoding="utf-8")

        existentes = sorted(pasta.glob("snapshot_*.json"))
        for antigo in existentes[:-MAX_SNAPSHOTS]:
            antigo.unlink(missing_ok=True)
        LOG.info("Snapshot gravado em %s", destino)
        return destino
    except Exception as erro:
        _registrar_erro("Falha ao gravar o snapshot de segurança", erro)
        return None


# ===========================================================================
# 4. EXPORTADORES
# ===========================================================================

LARGURA_UTIL_A4_PAISAGEM = landscape(A4)[0] - 30  # margens de 15pt


def _estilos_pdf() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle(
            "TituloOficial",
            parent=base["Heading1"],
            fontSize=12,
            alignment=1,
            spaceAfter=4,
        ),
        "subtitulo": ParagraphStyle(
            "SubtituloOficial",
            parent=base["Normal"],
            fontSize=9,
            alignment=1,
            spaceAfter=10,
        ),
        "grupo": ParagraphStyle(
            "GrupoLocalidades",
            parent=base["Heading2"],
            fontSize=10,
            textColor=colors.HexColor(COR_PRIMARIA),
            spaceBefore=10,
            spaceAfter=4,
        ),
        "local": ParagraphStyle(
            "NomeLocalidade",
            parent=base["Heading3"],
            fontSize=9,
            textColor=colors.HexColor(COR_SECUNDARIA),
            spaceBefore=6,
            spaceAfter=3,
        ),
        "nota": ParagraphStyle(
            "NotaRodape", parent=base["Normal"], fontSize=7.5, textColor=colors.grey
        ),
        "corpo": ParagraphStyle("Corpo", parent=base["Normal"], fontSize=8),
    }


def _cabecalho_oficial(elementos: list, estilos: dict, subtitulo: str) -> None:
    elementos.append(
        Paragraph(
            "<b>DETRAN-MA — DEPARTAMENTO ESTADUAL DE TRÂNSITO DO MARANHÃO</b>",
            estilos["titulo"],
        )
    )
    elementos.append(Paragraph(subtitulo, estilos["subtitulo"]))


def _estilo_tabela_padrao(linhas_destaque: Sequence[int] = ()) -> TableStyle:
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(COR_PRIMARIA)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
    ]
    for indice in linhas_destaque:
        estilo.extend(
            [
                ("BACKGROUND", (0, indice), (-1, indice), colors.HexColor(COR_CLARA)),
                ("FONTNAME", (0, indice), (-1, indice), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, indice), (-1, indice), colors.HexColor(COR_PRIMARIA)),
            ]
        )
    return TableStyle(estilo)


def _colunas_com_valor(df: pd.DataFrame, colunas: list[str]) -> list[str]:
    if df is None or df.empty:
        return []
    return [c for c in colunas if c in df.columns and df[c].fillna(0).sum() > 0]


def _num(valor: Any) -> int:
    return int(pd.to_numeric(valor, errors="coerce") or 0)


def _texto_ou_traco(valor: int) -> str:
    return "-" if valor == 0 else str(valor)


def _apenas_disponiveis(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()
    return df[(df["Status"] == STATUS_DISPONIVEL) & (df["Total"] > 0)].copy()


def _ordenar_por_data(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    copia = df.copy()
    copia["_ordem"] = pd.to_datetime(copia["Data"], format="%d/%m/%Y", errors="coerce")
    return copia.sort_values(["_ordem", "Horário"]).drop(columns=["_ordem"])


# --- PDF: grade detalhada (linha a linha por horário) -----------------------
COLUNAS_GRADE_DETALHADA = ["Data", "Dia da Semana", "Status", "Exam. M", "Exam. T", "Horário"]


def _linhas_grade_detalhada(
    df_filtrado: pd.DataFrame, pcd_usadas: list[str]
) -> tuple[list[str], list[list[str]], list[int]]:
    """Linha a linha por horário, com subtotal por data.

    Compartilhada pelo PDF de uma única localidade e pelo PDF completo da
    banca, para que os dois documentos tragam exatamente o mesmo nível de
    detalhe — é o formato que os colaboradores usam para lançar no sistema.
    Devolve (colunas, linhas, índices das linhas de subtotal a destacar).
    """
    numericas = COLS_CATEGORIA + pcd_usadas + ["Total"]
    colunas = COLUNAS_GRADE_DETALHADA + COLS_CATEGORIA + pcd_usadas + ["Total"]

    linhas: list[list[str]] = []
    destaques: list[int] = []
    for data_valor, grupo in df_filtrado.groupby("Data", sort=False):
        for _, registro in grupo.iterrows():
            linha = []
            for coluna in colunas:
                valor = registro.get(coluna, 0)
                if coluna in numericas:
                    linha.append(_texto_ou_traco(_num(valor)))
                else:
                    linha.append(html.escape(str(valor)))
            linhas.append(linha)

        subtotal = [f"TOTAL {data_valor}", "", "SUBTOTAL", "-", "-", "-"]
        for coluna in COLS_CATEGORIA + pcd_usadas:
            subtotal.append(_texto_ou_traco(int(grupo[coluna].fillna(0).sum())))
        subtotal.append(str(int(grupo["Total"].fillna(0).sum())))
        linhas.append(subtotal)
        destaques.append(len(linhas))

    return colunas, linhas, destaques


def gerar_pdf_localidade(
    df_dados: pd.DataFrame, banca: str, local: str, mes: str, ano: int | str
) -> bytes:
    """Grade de lançamento de vagas, linha a linha por horário."""
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=15,
        leftMargin=15,
        topMargin=15,
        bottomMargin=15,
        title=f"Grade de vagas — {local} — {mes}/{ano}",
    )
    estilos = _estilos_pdf()
    elementos: list = []
    _cabecalho_oficial(
        elementos,
        estilos,
        "GRADE DE LANÇAMENTO DE VAGAS DE EXAMES PRÁTICOS — "
        f"{html.escape(str(banca)).upper()} ({html.escape(str(local)).upper()})"
        f" — {html.escape(str(mes))}/{html.escape(str(ano))}",
    )

    df_filtrado = _ordenar_por_data(_apenas_disponiveis(df_dados))
    if df_filtrado.empty:
        elementos.append(
            Paragraph(
                "<b>Nenhuma vaga ofertada para os parâmetros selecionados.</b>",
                estilos["subtitulo"],
            )
        )
        documento.build(elementos)
        return buffer.getvalue()

    pcd_usadas = _colunas_com_valor(df_filtrado, COLS_PCD)
    colunas, linhas, destaques = _linhas_grade_detalhada(df_filtrado, pcd_usadas)

    tabela = Table([colunas] + linhas, repeatRows=1)
    tabela.setStyle(_estilo_tabela_padrao(destaques))
    elementos.append(tabela)
    documento.build(elementos)
    return buffer.getvalue()


# --- PDF: calendário consolidado da banca -----------------------------------


def gerar_pdf_banca(
    dados_por_local: dict[str, pd.DataFrame],
    viagens: Sequence[Viagem],
    bancas_config: dict[str, list[str]],
    banca: str,
    mes: str,
    ano: int,
    mes_numero: int,
    efetivo_maximo: int,
) -> bytes:
    """Calendário completo da banca, para conferência antes da publicação.

    As localidades saem na ordem pedida: unidades da sede, região
    metropolitana e depois os demais municípios pela data do primeiro exame.
    """
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=15,
        leftMargin=15,
        topMargin=15,
        bottomMargin=15,
        title=f"Calendário {banca} — {mes}/{ano}",
    )
    estilos = _estilos_pdf()
    elementos: list = []
    _cabecalho_oficial(
        elementos,
        estilos,
        "CALENDÁRIO CONSOLIDADO DE EXAMES PRÁTICOS — "
        f"{html.escape(str(banca)).upper()} — {html.escape(str(mes)).upper()}/{ano}",
    )

    ordenadas = ordenar_localidades_para_pdf(banca, bancas_config, dados_por_local)
    grupo_atual = None
    total_geral = 0
    houve_conteudo = False
    primeira_localidade = True

    for grupo, local in ordenadas:
        df_local_bruto = dados_por_local.get(local)
        df_local = _ordenar_por_data(_apenas_disponiveis(df_local_bruto))
        if df_local.empty:
            continue

        # Cada localidade começa em página nova: é o que permite entregar a
        # grade de um município isoladamente para o colaborador responsável
        # por lançar aquele calendário, sem recortar o PDF manualmente.
        if not primeira_localidade:
            elementos.append(PageBreak())
        primeira_localidade = False
        houve_conteudo = True

        if grupo != grupo_atual:
            elementos.append(Paragraph(html.escape(grupo), estilos["grupo"]))
            grupo_atual = grupo

        pcd_usadas = _colunas_com_valor(df_local, COLS_PCD)
        cabecalho, linhas, destaques = _linhas_grade_detalhada(df_local, pcd_usadas)

        total_local = int(df_local["Total"].fillna(0).sum())
        total_geral += total_local

        # Uma linha de total geral da localidade, além dos subtotais por
        # data: dá o resumo rápido sem perder o detalhe linha a linha.
        rodape = [f"TOTAL GERAL — {local}".upper(), "", "", "-", "-", "-"]
        for coluna in COLS_CATEGORIA + pcd_usadas:
            rodape.append(_texto_ou_traco(int(df_local[coluna].fillna(0).sum())))
        rodape.append(str(total_local))
        linhas.append(rodape)
        indice_total_geral = len(linhas)
        destaques_relativos = destaques + [indice_total_geral]

        dias_unicos = sorted(
            {d.day for d in (para_data(x) for x in df_local["Data"].unique()) if d}
        )
        legenda = (
            f"<b>{html.escape(local)}</b> — {len(dias_unicos)} dia(s) de exame:"
            f" {formatar_datas_exames(dias_unicos)} — {total_local} vagas"
        )

        tabela = Table(
            [cabecalho] + linhas,
            repeatRows=1,
            colWidths=_larguras_proporcionais(len(cabecalho)),
        )
        estilo_tabela = _estilo_tabela_padrao(destaques_relativos)
        # A linha de total geral se destaca mais que os subtotais por data.
        estilo_tabela.add(
            "BACKGROUND",
            (0, indice_total_geral),
            (-1, indice_total_geral),
            colors.HexColor(COR_SECUNDARIA),
        )
        estilo_tabela.add(
            "TEXTCOLOR", (0, indice_total_geral), (-1, indice_total_geral), colors.white
        )
        tabela.setStyle(estilo_tabela)
        elementos.append(Paragraph(legenda, estilos["local"]))
        elementos.append(tabela)

    if not houve_conteudo:
        elementos.append(
            Paragraph(
                "<b>Nenhuma vaga lançada nesta banca para o mês selecionado.</b>",
                estilos["subtitulo"],
            )
        )
        documento.build(elementos)
        return buffer.getvalue()

    # Painel de efetivo diário da banca no mês.
    elementos.append(PageBreak())
    elementos.append(
        Paragraph("CONFERÊNCIA DE EFETIVO DIÁRIO", estilos["grupo"])
    )
    tabela_efetivo = _tabela_efetivo_mensal(
        dados_por_local, viagens, banca, ano, mes_numero, efetivo_maximo
    )
    if tabela_efetivo:
        elementos.append(tabela_efetivo)
    elementos.append(Spacer(1, 6))
    elementos.append(
        Paragraph(
            f"Total geral de vagas da banca no mês: <b>{total_geral}</b>."
            " Documento gerado para conferência interna em "
            f"{datetime.datetime.now().strftime('%d/%m/%Y às %H:%M')}.",
            estilos["nota"],
        )
    )

    documento.build(elementos)
    return buffer.getvalue()


def _larguras_proporcionais(quantidade_colunas: int) -> list[float]:
    """Primeira coluna um pouco maior; o resto divide o espaço restante."""
    if quantidade_colunas <= 0:
        return []
    primeira = 52
    horarios = 62
    restante = LARGURA_UTIL_A4_PAISAGEM - primeira - horarios - 28
    demais = max(restante / max(quantidade_colunas - 3, 1), 26)
    larguras = [primeira, 28, horarios] + [demais] * (quantidade_colunas - 3)
    return larguras[:quantidade_colunas]


def _tabela_efetivo_mensal(
    dados_por_local: dict[str, pd.DataFrame],
    viagens: Sequence[Viagem],
    banca: str,
    ano: int,
    mes_numero: int,
    efetivo_maximo: int,
) -> Table | None:
    """Manhã, tarde e pico de examinadores por dia útil do mês."""
    totais: dict[datetime.date, dict[str, int]] = {}

    for local, df in dados_por_local.items():
        disponiveis = _apenas_disponiveis(df)
        if disponiveis.empty:
            continue
        for data_texto, grupo in disponiveis.groupby("Data"):
            data = para_data(data_texto)
            if not data:
                continue
            acumulado = totais.setdefault(data, {"M": 0, "T": 0})
            em_viagem_m = viagem_do_local(viagens, banca, local, data, TURNO_MANHA)
            em_viagem_t = viagem_do_local(viagens, banca, local, data, TURNO_TARDE)
            if not em_viagem_m:
                acumulado["M"] += _num(grupo["Exam. M"].max())
            if not em_viagem_t:
                acumulado["T"] += _num(grupo["Exam. T"].max())

    for numero_dia in range(1, calendar.monthrange(ano, mes_numero)[1] + 1):
        data = datetime.date(ano, mes_numero, numero_dia)
        manha, tarde = efetivo_em_viagem(viagens, banca, data)
        if manha or tarde:
            acumulado = totais.setdefault(data, {"M": 0, "T": 0})
            acumulado["M"] += manha
            acumulado["T"] += tarde

    if not totais:
        return None

    linhas = [["Data", "Dia", "Manhã", "Tarde", "Pico", "Limite", "Folga"]]
    destaques = []
    for data in sorted(totais):
        valores = totais[data]
        pico = pico_diario(valores["M"], valores["T"])
        linhas.append(
            [
                formatar_br(data),
                nome_dia_semana(data)[:3],
                str(valores["M"]),
                str(valores["T"]),
                str(pico),
                str(efetivo_maximo),
                str(efetivo_maximo - pico),
            ]
        )
        if pico > efetivo_maximo:
            destaques.append(len(linhas) - 1)

    tabela = Table(linhas, repeatRows=1)
    estilo = _estilo_tabela_padrao()
    for indice in destaques:
        estilo.add("BACKGROUND", (0, indice), (-1, indice), colors.HexColor("#FED7D7"))
        estilo.add("TEXTCOLOR", (0, indice), (-1, indice), colors.HexColor(COR_ALERTA))
    tabela.setStyle(estilo)
    return tabela


# --- PDF: resumo semanal do quadro de examinadores --------------------------
def gerar_pdf_semana(
    cabecalho_colunas: list[str],
    linhas_matriz: list[list[str]],
    equipes: list[dict[str, Any]],
    banca: str,
    mes: str,
    ano: int,
    rotulo_semana: str,
) -> bytes:
    """Escala semanal do quadro de examinadores, com as equipes em viagem."""
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=15,
        leftMargin=15,
        topMargin=15,
        bottomMargin=15,
        title=f"Escala semanal — {banca} — {rotulo_semana}",
    )
    estilos = _estilos_pdf()
    elementos: list = []
    _cabecalho_oficial(
        elementos,
        estilos,
        f"ESCALA SEMANAL DE EXAMINADORES — {html.escape(str(banca)).upper()}"
        f" — {html.escape(rotulo_semana).upper()} DE {html.escape(str(mes)).upper()}/{ano}",
    )

    dados = [cabecalho_colunas] + [
        [html.escape(str(celula)) for celula in linha] for linha in linhas_matriz
    ]
    largura_primeira = 150
    largura_demais = max(
        (LARGURA_UTIL_A4_PAISAGEM - largura_primeira)
        / max(len(cabecalho_colunas) - 1, 1),
        40,
    )
    tabela = Table(
        dados,
        repeatRows=1,
        colWidths=[largura_primeira]
        + [largura_demais] * (len(cabecalho_colunas) - 1),
    )
    estilo = _estilo_tabela_padrao([len(dados) - 1])
    estilo.add("ALIGN", (0, 1), (0, -1), "LEFT")
    tabela.setStyle(estilo)
    elementos.append(tabela)
    elementos.append(Spacer(1, 10))

    elementos.append(Paragraph("EQUIPES ITINERANTES NA SEMANA", estilos["grupo"]))
    if equipes:
        linhas_equipes = [
            ["Equipe", "Período", "Municípios", "Examinadores", "Observação"]
        ]
        for equipe in equipes:
            linhas_equipes.append(
                [
                    f"Banca {int(equipe['Numero']):02d}",
                    f"{formatar_br(equipe['Inicio Janela'])} a"
                    f" {formatar_br(equipe['Fim Janela'])}",
                    html.escape(" / ".join(equipe["Destinos"])),
                    f"{int(equipe['Examinadores']):02d}",
                    html.escape(equipe.get("Observacao", "") or "-"),
                ]
            )
        tabela_equipes = Table(
            linhas_equipes,
            repeatRows=1,
            colWidths=[60, 120, 240, 70, LARGURA_UTIL_A4_PAISAGEM - 490],
        )
        estilo_equipes = _estilo_tabela_padrao()
        estilo_equipes.add("ALIGN", (2, 1), (2, -1), "LEFT")
        estilo_equipes.add("ALIGN", (4, 1), (4, -1), "LEFT")
        tabela_equipes.setStyle(estilo_equipes)
        elementos.append(tabela_equipes)
    else:
        elementos.append(
            Paragraph("Nenhuma equipe em viagem nesta semana.", estilos["corpo"])
        )

    elementos.append(Spacer(1, 8))
    elementos.append(
        Paragraph(
            "Legenda: número = examinadores previstos · <b>Viagem</b> = equipe"
            " itinerante em deslocamento · <b>x</b> = sem lançamento ·"
            " <b>-</b> = dia fora da grade da localidade. Gerado em "
            f"{datetime.datetime.now().strftime('%d/%m/%Y às %H:%M')}.",
            estilos["nota"],
        )
    )
    documento.build(elementos)
    return buffer.getvalue()


# --- PDF: escala mensal de viagens itinerantes ------------------------------
def gerar_pdf_escala_viagens(
    viagens: Sequence[Viagem], mes: str, ano: int, mes_numero: int
) -> bytes:
    """Escala de viagens do mês, com preposto e veículo por equipe.

    A logística (01 preposto e 01 veículo por equipe itinerante) é atribuída
    automaticamente aqui: não há cadastro disso na interface, conforme
    solicitado — é informação que existe apenas neste documento.
    """
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=18,
        leftMargin=18,
        topMargin=18,
        bottomMargin=18,
        title=f"Escala de viagens — {mes}/{ano}",
    )
    estilos = _estilos_pdf()
    elementos: list = []
    _cabecalho_oficial(
        elementos,
        estilos,
        "ESCALA DE VIAGENS DAS BANCAS ITINERANTES — "
        f"{html.escape(str(mes)).upper()} DE {ano}",
    )

    primeiro = datetime.date(ano, mes_numero, 1)
    ultimo = datetime.date(ano, mes_numero, calendar.monthrange(ano, mes_numero)[1])

    grupos = agrupar_trechos_por_equipe(viagens)
    do_mes = []
    for grupo in grupos:
        dias = [
            d
            for t in grupo["Trechos"]
            for d in dias_efetivos_da_viagem(t)
            if primeiro <= d <= ultimo
        ]
        if not dias:
            continue
        copia = dict(grupo)
        copia["Inicio Mes"] = min(dias)
        copia["Fim Mes"] = max(dias)
        do_mes.append(copia)

    if not do_mes:
        elementos.append(
            Paragraph(
                "<b>Nenhuma viagem programada para o mês selecionado.</b>",
                estilos["subtitulo"],
            )
        )
        documento.build(elementos)
        return buffer.getvalue()

    total_examinadores_mes = 0
    total_prepostos = 0
    total_veiculos = 0

    for banca in sorted({g["Banca"] for g in do_mes}):
        elementos.append(Paragraph(html.escape(banca).upper(), estilos["grupo"]))

        linhas = [
            [
                "Banca\nItinerante",
                "Período",
                "Municípios",
                "Equipe",
                "Veículo",
                "Detalhamento por dia",
            ]
        ]
        for grupo in sorted(
            (g for g in do_mes if g["Banca"] == banca), key=lambda g: g["Numero"]
        ):
            examinadores = int(grupo["Examinadores"])
            total_examinadores_mes += examinadores
            total_prepostos += PREPOSTOS_POR_EQUIPE
            total_veiculos += VEICULOS_POR_EQUIPE

            detalhes = []
            for trecho in grupo["Trechos"]:
                dias = [
                    d
                    for d in dias_efetivos_da_viagem(trecho)
                    if primeiro <= d <= ultimo
                ]
                if not dias:
                    continue
                turnos = {turno_da_viagem_no_dia(trecho, d) for d in dias}
                if turnos == {TURNO_INTEGRAL}:
                    sufixo = ""
                else:
                    sufixo = " (" + ", ".join(
                        f"{formatar_curto(d)} {turno_da_viagem_no_dia(trecho, d).lower()}"
                        for d in dias
                        if turno_da_viagem_no_dia(trecho, d) != TURNO_INTEGRAL
                    ) + ")"
                detalhes.append(
                    f"{trecho.get('Destino')}: {formatar_curto(min(dias))} a"
                    f" {formatar_curto(max(dias))}{sufixo}"
                )

            observacoes = [
                str(t.get("Observações", "")).strip()
                for t in grupo["Trechos"]
                if str(t.get("Observações", "")).strip()
            ]
            if observacoes:
                detalhes.append("Obs.: " + "; ".join(dict.fromkeys(observacoes)))

            linhas.append(
                [
                    f"Banca {int(grupo['Numero']):02d}",
                    f"{formatar_br(grupo['Inicio Mes'])}\na"
                    f" {formatar_br(grupo['Fim Mes'])}",
                    Paragraph(
                        html.escape(" e ".join(grupo["Destinos"])), estilos["corpo"]
                    ),
                    f"{examinadores:02d} examinadores\n+"
                    f" {PREPOSTOS_POR_EQUIPE:02d} preposto",
                    f"{VEICULOS_POR_EQUIPE:02d} veículo",
                    Paragraph(html.escape(" · ".join(detalhes)), estilos["corpo"]),
                ]
            )

        tabela = Table(
            linhas,
            repeatRows=1,
            colWidths=[62, 92, 150, 96, 54, LARGURA_UTIL_A4_PAISAGEM - 454],
        )
        estilo = _estilo_tabela_padrao()
        estilo.add("ALIGN", (2, 1), (2, -1), "LEFT")
        estilo.add("ALIGN", (5, 1), (5, -1), "LEFT")
        estilo.add("VALIGN", (0, 1), (-1, -1), "TOP")
        tabela.setStyle(estilo)
        elementos.append(tabela)
        elementos.append(Spacer(1, 8))

    resumo = Table(
        [
            ["TOTAL DO MÊS", "Equipes", "Examinadores", "Prepostos", "Veículos"],
            [
                "",
                f"{len(do_mes):02d}",
                f"{total_examinadores_mes:02d}",
                f"{total_prepostos:02d}",
                f"{total_veiculos:02d}",
            ],
        ],
        colWidths=[140, 90, 110, 90, 90],
    )
    resumo.setStyle(_estilo_tabela_padrao([1]))
    elementos.append(resumo)
    elementos.append(Spacer(1, 6))
    elementos.append(
        Paragraph(
            "Cada banca itinerante conta com 01 preposto e 01 veículo atrelados à"
            " equipe de examinadores. Documento gerado em "
            f"{datetime.datetime.now().strftime('%d/%m/%Y às %H:%M')}.",
            estilos["nota"],
        )
    )
    documento.build(elementos)
    return buffer.getvalue()


# --- Excel de publicação ----------------------------------------------------
@st.cache_data(show_spinner=False)
def gerar_excel_publicacao(
    df_resumo: pd.DataFrame, mes: str, ano: int | str, incluir_pcd: bool = True
) -> bytes:
    """Planilha no layout usado para publicação do calendário mensal."""
    planilha = openpyxl.Workbook()
    aba = planilha.active
    aba.title = "Plan1"

    fonte = Font(name="Arial", size=10)
    centro = Alignment(horizontal="center", vertical="center")
    centro_quebra = Alignment(horizontal="center", vertical="center", wrap_text=True)
    esquerda = Alignment(horizontal="left", vertical="center")
    fina = Side(border_style="thin", color="000000")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)

    colunas_vagas = list(COLS_CATEGORIA)
    if incluir_pcd:
        colunas_vagas += _colunas_com_valor(df_resumo, COLS_PCD)

    larguras = {"A": 3, "B": 30.29}
    for indice in range(len(colunas_vagas)):
        larguras[chr(ord("C") + indice)] = 8.43
    larguras[chr(ord("C") + len(colunas_vagas))] = 25.29
    for letra, largura in larguras.items():
        aba.column_dimensions[letra].width = largura

    ultima_coluna = 2 + len(colunas_vagas) + 1
    linha_atual = 3

    for banca, grupo in df_resumo.groupby("Banca", sort=False):
        nome_publicacao = str(banca).replace("Banca ", "")
        aba.merge_cells(
            start_row=linha_atual,
            start_column=2,
            end_row=linha_atual,
            end_column=ultima_coluna,
        )
        titulo = aba.cell(
            row=linha_atual,
            column=2,
            value=(
                f"Quantidade de Exames Mensais de {nome_publicacao} e Região"
                f" - {str(mes).upper()} DE {ano}"
            ),
        )
        titulo.font = fonte
        titulo.alignment = centro
        for coluna in range(2, ultima_coluna + 1):
            aba.cell(row=linha_atual, column=coluna).border = borda
        linha_atual += 1

        for indice, texto in enumerate(
            ["Cidade"] + colunas_vagas + ["Datas dos exames"], 2
        ):
            celula = aba.cell(row=linha_atual, column=indice, value=texto)
            celula.font = fonte
            celula.alignment = centro
            celula.border = borda
        linha_atual += 1

        for _, registro in grupo.iterrows():
            celula_cidade = aba.cell(
                row=linha_atual, column=2, value=str(registro["Localidade"])
            )
            celula_cidade.font = fonte
            celula_cidade.alignment = esquerda
            celula_cidade.border = borda

            for indice, coluna in enumerate(colunas_vagas, 3):
                valor = _num(registro.get(coluna, 0))
                celula = aba.cell(
                    row=linha_atual, column=indice, value=valor if valor > 0 else None
                )
                celula.font = fonte
                celula.alignment = centro
                celula.number_format = "#,##0"
                celula.border = borda

            celula_datas = aba.cell(
                row=linha_atual,
                column=ultima_coluna,
                value=str(registro.get("Datas dos exames", "-")),
            )
            celula_datas.font = fonte
            celula_datas.alignment = centro_quebra
            celula_datas.number_format = "@"
            celula_datas.border = borda
            linha_atual += 1

        linha_atual += 1

    ultima_linha = max(linha_atual - 2, 3)
    aba.print_area = f"A3:{chr(ord('A') + ultima_coluna - 1)}{ultima_linha}"
    aba.page_setup.orientation = "landscape"
    aba.page_setup.paperSize = aba.PAPERSIZE_A4
    aba.page_margins.left = 0.51
    aba.page_margins.right = 0.51
    aba.page_margins.top = 0.79
    aba.page_margins.bottom = 0.79
    aba.page_margins.header = 0.31
    aba.page_margins.footer = 0.31

    buffer = io.BytesIO()
    planilha.save(buffer)
    return buffer.getvalue()


# ===========================================================================
# 5. COMPONENTES DE INTERFACE
# ===========================================================================


def _botao_abrir_pdf(pdf_bytes: bytes, chave: str, rotulo: str) -> None:
    """Botão que abre o PDF em uma nova aba do navegador.

    Usa Blob + URL.createObjectURL em vez de um link ``data:``: o Chrome
    bloqueia navegação de nível superior para URLs ``data:``, então o link
    simples não abriria a aba.
    """
    b64 = base64.b64encode(pdf_bytes).decode("ascii")
    identificador = re.sub(r"[^A-Za-z0-9_]", "_", chave)
    componente = f"""
    <button id="abrir_{identificador}" style="
        width: 100%;
        padding: 0.55rem 0.75rem;
        border-radius: 6px;
        border: 1px solid #2B6CB0;
        background: #2B6CB0;
        color: #fff;
        font-weight: 600;
        font-size: 0.9rem;
        font-family: 'Source Sans Pro', sans-serif;
        cursor: pointer;">{html.escape(rotulo)}</button>
    <script>
      (function() {{
        const dados = "{b64}";
        const botao = document.getElementById("abrir_{identificador}");
        botao.addEventListener("click", function() {{
          const binario = atob(dados);
          const bytes = new Uint8Array(binario.length);
          for (let i = 0; i < binario.length; i++) {{
            bytes[i] = binario.charCodeAt(i);
          }}
          const blob = new Blob([bytes], {{ type: "application/pdf" }});
          const url = URL.createObjectURL(blob);
          const aba = window.open(url, "_blank");
          if (!aba) {{
            alert("O navegador bloqueou a nova aba. Libere os pop-ups deste site.");
          }}
          setTimeout(function() {{ URL.revokeObjectURL(url); }}, 60000);
        }});
      }})();
    </script>
    """
    components.html(componente, height=52)


def _previa_pdf(pdf_bytes: bytes, chave: str, altura: int = 620) -> None:
    """Pré-visualização embutida do PDF, para conferir sem sair da página."""
    b64 = base64.b64encode(pdf_bytes).decode("ascii")
    identificador = re.sub(r"[^A-Za-z0-9_]", "_", chave)
    componente = f"""
    <div id="previa_{identificador}" style="width:100%;height:{altura}px;"></div>
    <script>
      (function() {{
        const dados = "{b64}";
        const binario = atob(dados);
        const bytes = new Uint8Array(binario.length);
        for (let i = 0; i < binario.length; i++) {{
          bytes[i] = binario.charCodeAt(i);
        }}
        const blob = new Blob([bytes], {{ type: "application/pdf" }});
        const url = URL.createObjectURL(blob);
        const quadro = document.createElement("iframe");
        quadro.src = url;
        quadro.style.width = "100%";
        quadro.style.height = "100%";
        quadro.style.border = "1px solid #CBD5E0";
        quadro.style.borderRadius = "6px";
        document.getElementById("previa_{identificador}").appendChild(quadro);
      }})();
    </script>
    """
    components.html(componente, height=altura + 10)


def bloco_pdf(
    *,
    chave: str,
    rotulo_gerar: str,
    gerador,
    nome_arquivo: str,
    ajuda: str | None = None,
    altura_previa: int = 620,
) -> None:
    """Fluxo padrão de PDF: gerar, conferir e só então baixar.

    O PDF é montado apenas quando solicitado — antes ele era remontado a cada
    interação com qualquer widget da página.
    """
    chave_bytes = f"_pdf_{chave}"
    chave_previa = f"_previa_{chave}"

    if st.button(rotulo_gerar, key=f"btn_gerar_{chave}", help=ajuda):
        try:
            st.session_state[chave_bytes] = gerador()
            st.session_state[chave_previa] = True
        except Exception as erro:
            LOG.exception("Falha ao gerar o PDF %s", chave)
            st.error(f"Não foi possível gerar o PDF: {erro}")

    pdf_bytes = st.session_state.get(chave_bytes)
    if not pdf_bytes:
        return

    tamanho_kb = len(pdf_bytes) / 1024
    st.caption(f"Documento pronto ({tamanho_kb:.0f} KB) — confira antes de baixar.")

    col_abrir, col_baixar = st.columns(2)
    with col_abrir:
        _botao_abrir_pdf(pdf_bytes, chave, "🔎 Abrir em nova aba")
    with col_baixar:
        st.download_button(
            "⬇️ Baixar PDF",
            data=pdf_bytes,
            file_name=nome_arquivo,
            mime="application/pdf",
            key=f"btn_baixar_{chave}",
            **LARGURA_TOTAL,
        )

    if st.checkbox(
        "👁️ Pré-visualizar aqui na página",
        value=bool(st.session_state.get(chave_previa)),
        key=f"chk_previa_{chave}",
    ):
        _previa_pdf(pdf_bytes, chave, altura_previa)


def dados_da_banca_no_mes(banca: str, mes: str, ano: int) -> dict[str, pd.DataFrame]:
    """Calendários já lançados de todas as localidades da banca no mês."""
    resultado: dict[str, pd.DataFrame] = {}
    for chave, df in st.session_state["historico_localidades"].items():
        partes = desmontar_chave_historico(chave)
        if not partes:
            continue
        banca_chave, mes_chave, ano_chave, local = partes
        if (banca_chave, mes_chave, ano_chave) == (banca, mes, str(ano)):
            resultado[local] = df
    return resultado


# ===========================================================================
# 6. ABAS
# ===========================================================================


# --- Aba: quadro matriz de examinadores -------------------------------------
def _agrupar_por_semana(
    dias: list[datetime.date],
) -> dict[datetime.date, list[datetime.date]]:
    """Agrupa dias úteis pela segunda-feira da semana (única, ao contrário do
    número ISO, que se repete na virada do ano)."""
    semanas: dict[datetime.date, list[datetime.date]] = {}
    for dia in dias:
        segunda = dia - datetime.timedelta(days=dia.weekday())
        semanas.setdefault(segunda, []).append(dia)
    return dict(sorted(semanas.items()))


def _celula_do_quadro(
    banca: str,
    local: str,
    dia: datetime.date,
    df_local: pd.DataFrame | None,
    feriados: frozenset[datetime.date],
    viagens: Sequence[Viagem],
) -> str:
    """Conteúdo da célula do quadro matriz.

    Localidade atendida por equipe itinerante mostra apenas a marcação de
    viagem: o efetivo dela é contabilizado uma única vez na linha de total,
    pela equipe, e não somado por município.
    """
    viagem = viagem_do_local_no_dia(viagens, banca, local, dia)
    if viagem:
        numero = int(viagem.get("Numero Banca Itinerante", 0) or 0)
        turno = turno_da_viagem_no_dia(viagem, dia)
        marcador = f"Viagem {numero:02d}"
        if turno in (TURNO_MANHA, TURNO_TARDE):
            marcador += " (M)" if turno == TURNO_MANHA else " (T)"
        return marcador

    data_texto = dia.strftime("%d/%m/%Y")
    if df_local is not None and not df_local.empty:
        do_dia = df_local[df_local["Data"] == data_texto]
        if not do_dia.empty and (do_dia["Status"] == STATUS_FERIADO).any():
            return "Feriado"
        disponiveis = do_dia[do_dia["Status"] == STATUS_DISPONIVEL]
        if not disponiveis.empty:
            return str(
                pico_diario(
                    _num(disponiveis["Exam. M"].max()),
                    _num(disponiveis["Exam. T"].max()),
                )
            )

    if dia in feriados:
        return "Feriado"
    return "x"


def _montar_matriz_semana(
    banca: str,
    mes_nome: str,
    ano: int,
    dias_semana: list[datetime.date],
    feriados: frozenset[datetime.date],
) -> tuple[list[str], list[list[str]]]:
    bancas_config = st.session_state["bancas_config"]
    viagens = st.session_state["viagens_registradas"]
    historico = st.session_state["historico_localidades"]
    dias_permitidos_dict = st.session_state["dias_permitidos_dict"]
    localidades = bancas_config.get(banca, [])

    linhas: list[list[str]] = []
    for local in localidades:
        linha = [local]
        df_local = historico.get(chave_historico(banca, mes_nome, ano, local))
        permitidos = dias_permitidos_dict.get(
            chave_local(banca, local), DIAS_UTEIS_PADRAO
        )
        for dia in dias_semana:
            if nome_dia_semana(dia) not in permitidos and not viagem_do_local_no_dia(
                viagens, banca, local, dia
            ):
                linha.append("-")
                continue
            linha.append(
                _celula_do_quadro(banca, local, dia, df_local, feriados, viagens)
            )
        linhas.append(linha)

    linha_total = ["TOTAL EXAMINADORES EM CAMPO"]
    for dia in dias_semana:
        total_manha = 0
        total_tarde = 0
        for local in localidades:
            permitidos = dias_permitidos_dict.get(
                chave_local(banca, local), DIAS_UTEIS_PADRAO
            )
            if nome_dia_semana(dia) not in permitidos:
                continue
            df_local = historico.get(chave_historico(banca, mes_nome, ano, local))
            if df_local is None or df_local.empty:
                continue
            do_dia = df_local[
                (df_local["Data"] == dia.strftime("%d/%m/%Y"))
                & (df_local["Status"] == STATUS_DISPONIVEL)
            ]
            if do_dia.empty:
                continue
            # Localidade em viagem não soma pelo município: o efetivo entra
            # uma única vez pela equipe itinerante, logo abaixo.
            if not viagem_do_local(viagens, banca, local, dia, TURNO_MANHA):
                total_manha += _num(do_dia["Exam. M"].max())
            if not viagem_do_local(viagens, banca, local, dia, TURNO_TARDE):
                total_tarde += _num(do_dia["Exam. T"].max())

        manha_viagem, tarde_viagem = efetivo_em_viagem(viagens, banca, dia)
        total_manha += manha_viagem
        total_tarde += tarde_viagem
        linha_total.append(str(pico_diario(total_manha, total_tarde)))
    linhas.append(linha_total)

    cabecalho = ["Localidade"] + [d.strftime("%d/%m (%a)") for d in dias_semana]
    return cabecalho, linhas


def aba_quadro() -> None:
    st.markdown("### 🗓️ Quadro Matriz de Distribuição de Examinadores")

    bancas_config = st.session_state["bancas_config"]
    disponiveis = [b for b in BANCAS_COM_QUADRO_MATRIZ if b in bancas_config]
    if not disponiveis:
        st.info("Nenhuma das bancas com quadro matriz está cadastrada no momento.")
        return

    col_banca, col_mes, col_ano = st.columns(3)
    banca = col_banca.selectbox("Selecione a Banca:", disponiveis, key="q_banca")
    mes_nome = col_mes.selectbox(
        "Mês de Visualização:", MESES_LISTA, index=indice_mes_atual(), key="q_mes"
    )
    ano = col_ano.selectbox(
        "Ano de Visualização:", anos_disponiveis(), index=0, key="q_ano"
    )
    mes_numero = numero_do_mes(mes_nome)

    st.markdown("---")

    calendario_mes = calendar.Calendar(firstweekday=0)
    dias_uteis = [
        d
        for d in calendario_mes.itermonthdates(ano, mes_numero)
        if d.month == mes_numero and d.weekday() < 5
    ]
    if not dias_uteis:
        st.info("Não há dias úteis neste mês.")
        return
    if not bancas_config.get(banca):
        st.warning(f"A banca {banca} não possui localidades cadastradas.")
        return

    feriados = feriados_estaduais(ano)
    viagens = st.session_state["viagens_registradas"]

    for indice, (_segunda, dias_semana) in enumerate(
        _agrupar_por_semana(dias_uteis).items(), 1
    ):
        rotulo = (
            f"Semana {indice} ({dias_semana[0].strftime('%d')} a"
            f" {dias_semana[-1].strftime('%d')} de {mes_nome})"
        )
        st.markdown(f"#### 📅 {rotulo}")

        cabecalho, linhas = _montar_matriz_semana(
            banca, mes_nome, ano, dias_semana, feriados
        )
        st.dataframe(
            pd.DataFrame(linhas, columns=cabecalho),
            **LARGURA_TOTAL,
            hide_index=True,
        )

        equipes = equipes_ativas_no_intervalo(
            viagens, banca, dias_semana[0], dias_semana[-1]
        )
        total_em_viagem = sum(int(e["Examinadores"]) for e in equipes)

        if equipes:
            st.caption(
                f"🚍 {len(equipes)} equipe(s) itinerante(s) nesta semana —"
                f" {total_em_viagem} examinadores em deslocamento."
            )
            for equipe in equipes:
                titulo = (
                    f"🚍 Banca {int(equipe['Numero']):02d} — Viagem para"
                    f" {' / '.join(equipe['Destinos'])}"
                    f" ({int(equipe['Examinadores'])} examinadores)"
                )
                with st.expander(titulo, expanded=False):
                    for trecho in equipe["Trechos"]:
                        dias_trecho = [
                            d
                            for d in dias_efetivos_da_viagem(trecho)
                            if dias_semana[0] <= d <= dias_semana[-1]
                        ]
                        if not dias_trecho:
                            continue
                        detalhe_turnos = descricao_turnos(trecho)
                        st.write(
                            f"• **{trecho.get('Destino')}** — "
                            f"{formatar_br(min(dias_trecho))} a"
                            f" {formatar_br(max(dias_trecho))} — "
                            f"{trecho.get('Turno', TURNO_INTEGRAL)}"
                            f" · {examinadores_da_banca_na_viagem(trecho, banca)}"
                            " examinador(es)"
                            + (f" · ajustes: {detalhe_turnos}" if detalhe_turnos else "")
                        )
                        if trecho.get("Observações"):
                            st.caption(str(trecho["Observações"]))
        else:
            st.caption("Nenhuma equipe itinerante nesta semana.")

        for equipe in equipes:
            equipe["Observacao"] = descricao_turnos(equipe["Trechos"][0])

        bloco_pdf(
            chave=f"semana_{banca}_{mes_nome}_{ano}_{indice}",
            rotulo_gerar=f"📄 Gerar PDF da Semana {indice}",
            gerador=lambda cab=cabecalho, lin=linhas, eq=equipes, rot=rotulo: (
                gerar_pdf_semana(cab, lin, eq, banca, mes_nome, ano, rot)
            ),
            nome_arquivo=(
                f"Escala_Semana{indice}_{banca}_{mes_nome}_{ano}.pdf".replace(" ", "_")
            ),
            ajuda="Gera o resumo da escala desta semana para impressão.",
            altura_previa=520,
        )
        st.markdown("---")

    st.caption(
        "Legenda: número = examinadores previstos · **Viagem NN** = equipe"
        " itinerante (M = só manhã, T = só tarde) · **x** = sem lançamento ·"
        " **-** = dia fora da grade da localidade."
    )


# --- Aba: montagem do calendário detalhado ----------------------------------
def _montar_grade_base(
    *,
    banca: str,
    local: str,
    ano: int,
    mes_numero: int,
    dias_permitidos: list[str],
    horarios: list[str],
    feriados: set[datetime.date],
    padrao_manha: int,
    padrao_tarde: int,
) -> tuple[pd.DataFrame, list[str]]:
    """Gera uma linha por (dia de atendimento × horário ativo)."""
    viagens = st.session_state["viagens_registradas"]
    bancas_config = st.session_state["bancas_config"]

    viagens_do_local = [
        v
        for v in viagens
        if banca_vinculada(v, banca) and v.get("Destino") == local
    ]
    eh_regional = banca not in ("Banca São Luís", "Banca Imperatriz")
    eh_sede = local == obter_sede(bancas_config, banca)

    total_dias = calendar.monthrange(ano, mes_numero)[1]
    registros: list[dict] = []
    datas_do_mes: list[str] = []

    for numero_dia in range(1, total_dias + 1):
        data = datetime.date(ano, mes_numero, numero_dia)
        nome_dia = DIAS_SEMANA_OPCOES[data.weekday()]
        tem_viagem_no_dia = any(viagem_cobre_dia(v, data) for v in viagens_do_local)
        if nome_dia not in dias_permitidos and not tem_viagem_no_dia:
            continue

        data_texto = data.strftime("%d/%m/%Y")
        if data_texto not in datas_do_mes:
            datas_do_mes.append(data_texto)
        eh_feriado = data in feriados

        for horario in horarios:
            periodo = periodo_do_horario(horario)
            viagem_no_periodo = next(
                (v for v in viagens_do_local if viagem_cobre_periodo(v, data, periodo)),
                None,
            )
            banca_viajando = viagens_ativas_no_periodo(viagens, banca, data, periodo)

            if eh_feriado:
                status, exam_manha, exam_tarde = STATUS_FERIADO, 0, 0
            elif viagem_no_periodo:
                efetivo = examinadores_da_banca_na_viagem(viagem_no_periodo, banca)
                status = STATUS_DISPONIVEL
                exam_manha = efetivo if periodo == TURNO_MANHA else 0
                exam_tarde = efetivo if periodo == TURNO_TARDE else 0
            elif tem_viagem_no_dia:
                # A equipe está no município, mas não neste turno.
                status, exam_manha, exam_tarde = STATUS_INDISPONIVEL, 0, 0
            elif eh_regional and eh_sede and banca_viajando:
                status, exam_manha, exam_tarde = STATUS_INDISPONIVEL, 0, 0
            elif nome_dia not in dias_permitidos:
                status, exam_manha, exam_tarde = STATUS_INDISPONIVEL, 0, 0
            else:
                status = STATUS_DISPONIVEL
                exam_manha = padrao_manha if periodo == TURNO_MANHA else 0
                exam_tarde = padrao_tarde if periodo == TURNO_TARDE else 0

            registro = {
                "Data": data_texto,
                "Dia da Semana": nome_dia,
                "Status": status,
                "Exam. M": exam_manha if status == STATUS_DISPONIVEL else 0,
                "Exam. T": exam_tarde if status == STATUS_DISPONIVEL else 0,
                "Horário": horario,
            }
            for coluna in COLS_VAGAS:
                registro[coluna] = 0
            registro["Total"] = 0
            registros.append(registro)

    colunas = (
        ["Data", "Dia da Semana", "Status", "Exam. M", "Exam. T", "Horário"]
        + COLS_VAGAS
        + ["Total"]
    )
    return pd.DataFrame(registros, columns=colunas), datas_do_mes


def _mesclar_com_historico(df_base: pd.DataFrame, chave: str) -> pd.DataFrame:
    """Preserva os lançamentos já feitos, sem deixar escapar bloqueios novos."""
    historico = st.session_state["historico_localidades"]
    if chave not in historico or df_base.empty:
        return df_base.copy()

    df_antigo = historico[chave]
    if df_antigo.empty:
        return df_base.copy()

    # Sem isto, um horário repetido faz o .loc devolver DataFrame em vez de Series.
    df_antigo = df_antigo.drop_duplicates(subset=["Data", "Horário"], keep="last")
    indexado = df_antigo.set_index(["Data", "Horário"])

    linhas = []
    for registro in df_base.to_dict(orient="records"):
        chave_linha = (registro["Data"], registro["Horário"])
        if registro["Status"] in STATUS_BLOQUEADOS or chave_linha not in indexado.index:
            linhas.append(registro)
            continue
        antigo = indexado.loc[chave_linha]
        registro["Status"] = antigo.get("Status", registro["Status"])
        registro["Exam. M"] = _num(antigo.get("Exam. M", 0))
        registro["Exam. T"] = _num(antigo.get("Exam. T", 0))
        for coluna in COLS_VAGAS:
            registro[coluna] = _num(antigo.get(coluna, 0))
        # Sem isto, o Total ficava zerado (herdado de df_base) até o editor
        # rodar de novo — o alerta de deslocamento e o destaque verde (Ajustes
        # 2 e 3) liam esse valor obsoleto assim que a página abria.
        registro["Total"] = _num(antigo.get("Total", 0))
        linhas.append(registro)
    return pd.DataFrame(linhas, columns=df_base.columns)


# Colunas desabilitadas no editor: são as únicas onde o Streamlit realmente
# aplica a cor de fundo do Styler (a documentação do data_editor é explícita
# quanto a isso — estilo em coluna editável é ignorado).
_COLUNAS_DESTACAVEIS_EDITOR = ["Data", "Dia da Semana", "Horário", "Total"]


def _cor_linha_calendario(linha: pd.Series, horario_minimo_min: int | None) -> str:
    """Verde: já lançada. Vermelho: lançada antes do horário sugerido pelo
    deslocamento (Ajuste 3). Cinza: bloqueada (feriado/indisponível)."""
    if linha.get("Status") in STATUS_BLOQUEADOS:
        return f"background-color: {COR_INATIVO_BG}; color: {COR_INATIVO_TXT};"
    lancada = _num(linha.get("Total")) > 0
    if lancada and horario_minimo_min is not None:
        minutos_linha = _hhmm_para_minutos(str(linha.get("Horário", "")))
        if minutos_linha is not None and minutos_linha < horario_minimo_min:
            return (
                f"background-color: {COR_ALERTA_DESLOC_BG};"
                f" color: {COR_ALERTA_DESLOC_TXT}; font-weight: 600;"
            )
    if lancada:
        return (
            f"background-color: {COR_LANCADO_BG}; color: {COR_LANCADO_TXT};"
            " font-weight: 600;"
        )
    return ""


def _estilizar_grade_calendario(df: pd.DataFrame, horario_minimo_min: int | None = None):
    """Aplica o destaque visual das linhas já preenchidas (Ajuste 2) e o
    alerta de deslocamento (Ajuste 3).

    O Streamlit só renderiza o estilo do pandas.Styler nas colunas
    desabilitadas do data_editor; nas colunas editáveis (Status, Exam. M/T,
    Cat A-E, PCD A-E) o estilo é ignorado silenciosamente. Por isso a cor é
    aplicada só em Data/Dia da Semana/Horário/Total — o suficiente para criar
    uma faixa verde bem visível em cada linha já lançada, sem quebrar a
    edição das demais colunas.
    """
    if df.empty:
        return df

    def _aplicar(linha: pd.Series) -> list[str]:
        estilo = _cor_linha_calendario(linha, horario_minimo_min)
        return [estilo if col in _COLUNAS_DESTACAVEIS_EDITOR else "" for col in df.columns]

    return df.style.apply(_aplicar, axis=1)


def _calendario_visual(
    banca: str,
    local: str,
    ano: int,
    mes_numero: int,
    feriados: set[datetime.date],
    chave: str,
) -> None:
    st.markdown("### 📅 Visualização em Calendário Interativo")
    if not TEM_COMPONENTE_CALENDARIO:
        st.info(
            "💡 Para arrastar viagens diretamente na tela, adicione"
            " `streamlit-calendar` ao `requirements.txt`."
        )
        return

    eventos = []
    for viagem in st.session_state["viagens_registradas"]:
        if not banca_vinculada(viagem, banca):
            continue
        inicio = para_data(viagem["Data Inicio"])
        fim = para_data(viagem["Data Fim"])
        if not inicio or not fim:
            continue
        eventos.append(
            {
                # O id viaja no evento: dispensa extrair o destino do título.
                "id": str(viagem.get("id") or ""),
                "title": (
                    f"🚍 {viagem['Destino']}"
                    f" ({examinadores_da_banca_na_viagem(viagem, banca)} Ex.)"
                ),
                "start": inicio.isoformat(),
                "end": (fim + datetime.timedelta(days=1)).isoformat(),
                "color": COR_SECUNDARIA if viagem["Destino"] == local else COR_NEUTRA,
                "allDay": True,
            }
        )

    for feriado in feriados:
        if feriado.month == mes_numero and feriado.year == ano:
            eventos.append(
                {
                    "id": "",
                    "title": "🔴 Feriado / Sem Atendimento",
                    "start": feriado.isoformat(),
                    "color": COR_ALERTA,
                    "allDay": True,
                }
            )

    opcoes = {
        "editable": True,
        "selectable": False,
        "headerToolbar": {
            "left": "prev,next today",
            "center": "title",
            "right": "dayGridMonth,timeGridWeek",
        },
        "initialDate": f"{ano}-{mes_numero:02d}-01",
        "locale": "pt-br",
    }

    retorno = componente_calendario(
        events=eventos, options=opcoes, key=f"cal_visual_{chave}"
    )
    alteracao = (retorno or {}).get("eventChange")
    if not alteracao:
        return

    evento = (alteracao or {}).get("event") or {}
    id_viagem = evento.get("id")
    if not id_viagem:
        return
    try:
        novo_inicio = datetime.date.fromisoformat(str(evento["start"])[:10])
        novo_fim = datetime.date.fromisoformat(
            str(evento["end"])[:10]
        ) - datetime.timedelta(days=1)
    except (KeyError, TypeError, ValueError):
        LOG.warning("Evento de calendário com datas ilegíveis: %r", evento)
        return
    if novo_fim < novo_inicio:
        novo_fim = novo_inicio

    for viagem in st.session_state["viagens_registradas"]:
        if str(viagem.get("id")) == str(id_viagem):
            viagem["Data Inicio"] = novo_inicio
            viagem["Data Fim"] = novo_fim
            # Ajustes de turno que caíram fora do novo intervalo perdem sentido.
            viagem["Turnos Por Data"] = {
                iso: turno
                for iso, turno in (viagem.get("Turnos Por Data") or {}).items()
                if (d := para_data(iso)) and novo_inicio <= d <= novo_fim
            }
            salvar_viagens(st.session_state["viagens_registradas"])
            st.toast(f"Viagem para {viagem['Destino']} reagendada.")
            st.rerun()


def _monitor_efetivo(
    banca: str, mes_nome: str, ano: int, datas_do_mes: list[str], efetivo_maximo: int
) -> None:
    st.sidebar.markdown("---")
    st.sidebar.subheader("🔍 Consulta de Efetivo por Data")
    if not datas_do_mes:
        st.sidebar.info("Nenhuma data de atendimento neste mês.")
        return

    data_consulta = st.sidebar.selectbox(
        "Escolha a Data para checar o total:", datas_do_mes, key="cal_data_consulta"
    )
    data = para_data(data_consulta)
    if not data:
        return

    viagens = st.session_state["viagens_registradas"]
    detalhes: list[dict] = []
    total_manha = 0
    total_tarde = 0

    for local, df_local in dados_da_banca_no_mes(banca, mes_nome, ano).items():
        do_dia = df_local[
            (df_local["Data"] == data_consulta)
            & (df_local["Status"] == STATUS_DISPONIVEL)
        ]
        if do_dia.empty:
            continue
        manha = (
            0
            if viagem_do_local(viagens, banca, local, data, TURNO_MANHA)
            else _num(do_dia["Exam. M"].max())
        )
        tarde = (
            0
            if viagem_do_local(viagens, banca, local, data, TURNO_TARDE)
            else _num(do_dia["Exam. T"].max())
        )
        total_manha += manha
        total_tarde += tarde
        if manha or tarde:
            detalhes.append({"Localidade": local, "Manhã": manha, "Tarde": tarde})

    equipes = equipes_em_viagem(viagens, banca, data)
    total_manha += sum(equipes[TURNO_MANHA].values())
    total_tarde += sum(equipes[TURNO_TARDE].values())
    for equipe in sorted(set(equipes[TURNO_MANHA]) | set(equipes[TURNO_TARDE])):
        detalhes.append(
            {
                "Localidade": f"Banca itinerante {equipe:02d} (Viagem)",
                "Manhã": equipes[TURNO_MANHA].get(equipe, 0),
                "Tarde": equipes[TURNO_TARDE].get(equipe, 0),
            }
        )

    pico = pico_diario(total_manha, total_tarde)
    st.sidebar.markdown(f"**Escala do Dia:** `{data_consulta}`")
    st.sidebar.metric("Efetivo Usado no Dia", f"{pico} / {efetivo_maximo}")
    st.sidebar.metric("Saldo Restante", efetivo_maximo - pico)
    if pico > efetivo_maximo:
        st.sidebar.error(
            f"⚠️ **ESTOURO DE EFETIVO!** Excesso de {pico - efetivo_maximo}"
            " examinadores."
        )
    if detalhes:
        st.sidebar.markdown("**Distribuição por Localidade nesta Data:**")
        st.sidebar.dataframe(
            pd.DataFrame(detalhes), **LARGURA_TOTAL, hide_index=True
        )
    else:
        st.sidebar.info("Sem exames ou viagens nesta data.")


def _pico_do_mes(banca: str, mes_nome: str, ano: int, mes_numero: int) -> int:
    totais: dict[str, dict[str, int]] = {}
    viagens = st.session_state["viagens_registradas"]

    for local, df_local in dados_da_banca_no_mes(banca, mes_nome, ano).items():
        disponiveis = df_local[df_local["Status"] == STATUS_DISPONIVEL]
        if disponiveis.empty:
            continue
        for data_texto, grupo in disponiveis.groupby("Data"):
            data = para_data(data_texto)
            if not data:
                continue
            acumulado = totais.setdefault(data_texto, {"M": 0, "T": 0})
            if not viagem_do_local(viagens, banca, local, data, TURNO_MANHA):
                acumulado["M"] += _num(grupo["Exam. M"].max())
            if not viagem_do_local(viagens, banca, local, data, TURNO_TARDE):
                acumulado["T"] += _num(grupo["Exam. T"].max())

    for numero_dia in range(1, calendar.monthrange(ano, mes_numero)[1] + 1):
        data = datetime.date(ano, mes_numero, numero_dia)
        manha, tarde = efetivo_em_viagem(viagens, banca, data)
        if manha or tarde:
            acumulado = totais.setdefault(data.strftime("%d/%m/%Y"), {"M": 0, "T": 0})
            acumulado["M"] += manha
            acumulado["T"] += tarde

    if not totais:
        return 0
    return max(pico_diario(v["M"], v["T"]) for v in totais.values())


def aba_calendario() -> None:
    bancas_config = st.session_state["bancas_config"]
    if not bancas_config:
        st.warning("Cadastre uma banca na aba de Gestão de Localidades.")
        return

    st.sidebar.header("⚙️ Parâmetros do Calendário")
    banca = st.sidebar.selectbox(
        "Banca Principal:", list(bancas_config.keys()), key="cal_banca"
    )
    localidades = bancas_config.get(banca, [])
    if not localidades:
        st.warning(f"A banca **{banca}** ainda não possui localidades cadastradas.")
        return
    local = st.sidebar.selectbox("Local de Atendimento:", localidades, key="cal_local")

    ano = st.sidebar.selectbox("Ano:", anos_disponiveis(), index=0, key="cal_ano")
    mes_nome = st.sidebar.selectbox(
        "Mês:", MESES_LISTA, index=indice_mes_atual(), key="cal_mes"
    )
    mes_numero = numero_do_mes(mes_nome)

    st.sidebar.markdown("---")
    st.sidebar.subheader("🛡️ Efetivo Diário da Banca")
    efetivo_maximo = st.sidebar.number_input(
        f"Efetivo Máximo Diário ({banca}):",
        min_value=1,
        max_value=200,
        value=int(st.session_state["capacidade_bancas"].get(banca, 30)) or 30,
        key="cal_efetivo",
    )

    st.sidebar.subheader("📅 Dias de Atendimento na Semana")
    chave_loc = chave_local(banca, local)
    dias_permitidos = st.sidebar.multiselect(
        "Selecione os dias de exame desta localidade:",
        options=DIAS_SEMANA_OPCOES,
        default=st.session_state["dias_permitidos_dict"].get(
            chave_loc, DIAS_UTEIS_PADRAO
        ),
        key=f"cal_dias_{chave_loc}",
    )
    st.session_state["dias_permitidos_dict"][chave_loc] = dias_permitidos
    salvar_dias_permitidos(st.session_state["dias_permitidos_dict"])

    st.sidebar.subheader("👨‍⚖️ Sugestão Inicial de Examinadores")
    padrao_manha = st.sidebar.number_input(
        "Padrão Inicial Manhã:", min_value=0, max_value=50, value=4, key="cal_pad_m"
    )
    padrao_tarde = st.sidebar.number_input(
        "Padrão Inicial Tarde:", min_value=0, max_value=50, value=4, key="cal_pad_t"
    )

    st.sidebar.subheader("🏙️ Feriados Municipais / Locais")
    texto_feriados = st.sidebar.text_area(
        "Digite um por linha (Ex: DD/MM - Motivo):",
        value=st.session_state["feriados_locais_dict"].get(chave_loc, ""),
        help="Formatos aceitos: DD/MM, DD/MM/AAAA ou AAAA-MM-DD",
        key=f"cal_feriados_{chave_loc}",
    )
    st.session_state["feriados_locais_dict"][chave_loc] = texto_feriados
    salvar_feriados_locais(st.session_state["feriados_locais_dict"])

    feriados = set(feriados_estaduais(ano))
    feriados |= datas_de_feriado_do_texto(texto_feriados, ano)

    horarios_ativos = [
        h
        for h in st.session_state["lista_horarios"]
        if h not in st.session_state["horarios_inativos"]
    ]
    if not horarios_ativos:
        st.warning(
            "Não há horários ativos. Cadastre ou reative horários na aba"
            " **Horários & Turmas**."
        )
        return
    if not dias_permitidos:
        st.warning("Selecione ao menos um dia da semana na barra lateral.")
        return

    df_base, datas_do_mes = _montar_grade_base(
        banca=banca,
        local=local,
        ano=ano,
        mes_numero=mes_numero,
        dias_permitidos=dias_permitidos,
        horarios=horarios_ativos,
        feriados=feriados,
        padrao_manha=padrao_manha,
        padrao_tarde=padrao_tarde,
    )

    chave = chave_historico(banca, mes_nome, ano, local)
    df_completo = _mesclar_com_historico(df_base, chave)

    _calendario_visual(banca, local, ano, mes_numero, feriados, chave)

    st.markdown("### 📝 Lançamento e Edição de Vagas por Horário")

    viagens_do_local = [
        v
        for v in st.session_state["viagens_registradas"]
        if banca_vinculada(v, banca) and v.get("Destino") == local
    ]
    if viagens_do_local:
        referencia = viagens_do_local[0]
        ajustes = descricao_turnos(referencia)
        st.success(
            f"🚍 **Viagem detectada**: {local} tem viagem registrada de"
            f" **{formatar_br(para_data(referencia['Data Inicio']))}** a"
            f" **{formatar_br(para_data(referencia['Data Fim']))}** com"
            f" **{examinadores_da_banca_na_viagem(referencia, banca)} examinadores**"
            f" ({referencia.get('Turno', TURNO_INTEGRAL)})."
            + (f" Ajustes por dia: {ajustes}." if ajustes else "")
        )
    else:
        st.info(f"📍 Configurando calendário para **{local}** ({banca}).")

    if df_completo.empty:
        st.warning("Nenhuma data de atendimento gerada para os parâmetros atuais.")
        return

    total_dias = calendar.monthrange(ano, mes_numero)[1]
    primeiro_dia = datetime.date(ano, mes_numero, 1)
    ultimo_dia = datetime.date(ano, mes_numero, total_dias)

    col_periodo, col_datas = st.columns([1, 2])
    with col_periodo:
        periodo = st.date_input(
            "🔍 Filtrar por Período de Atendimento:",
            value=(primeiro_dia, ultimo_dia),
            min_value=primeiro_dia,
            max_value=ultimo_dia,
            format="DD/MM/YYYY",
            key=f"cal_filtro_periodo_{chave}",
        )
    with col_datas:
        datas_selecionadas = st.multiselect(
            "📌 Ou selecione dias específicos do mês:",
            options=datas_do_mes,
            default=[],
            placeholder="Selecione um ou mais dias...",
            key=f"cal_filtro_datas_{chave}",
        )

    if datas_selecionadas:
        df_exibicao = df_completo[df_completo["Data"].isin(datas_selecionadas)].copy()
    elif isinstance(periodo, (tuple, list)) and len(periodo) == 2:
        inicio, fim = periodo
        convertidas = pd.to_datetime(df_completo["Data"], format="%d/%m/%Y").dt.date
        df_exibicao = df_completo[
            (convertidas >= inicio) & (convertidas <= fim)
        ].copy()
    else:
        df_exibicao = df_completo.copy()

    # Ajuste 3 — alerta de deslocamento: linhas com vagas lançadas antes do
    # horário em que a banca estimadamente chega ao município.
    minutos_desloc = tempo_deslocamento_efetivo(banca, local)
    horario_minimo_min = None
    linhas_em_alerta = pd.DataFrame()
    if minutos_desloc > 0:
        chegada_estimada = horario_chegada_estimado(minutos_desloc)
        horario_minimo_min = chegada_estimada.hour * 60 + chegada_estimada.minute
        minutos_das_linhas = df_exibicao["Horário"].apply(_hhmm_para_minutos)
        linhas_em_alerta = df_exibicao[
            (df_exibicao["Total"] > 0)
            & (df_exibicao["Status"] == STATUS_DISPONIVEL)
            & minutos_das_linhas.notna()
            & (minutos_das_linhas < horario_minimo_min)
        ]

    if not linhas_em_alerta.empty:
        chave_ignorar = f"ignorar_alerta_desloc_{chave}"
        if not st.session_state.get(chave_ignorar):
            datas_em_alerta = ", ".join(
                sorted(set(linhas_em_alerta["Data"]), key=lambda d: para_data(d) or datetime.date.max)
            )
            st.warning(
                f"⚠️ {len(linhas_em_alerta)} horário(s) com vagas lançadas antes"
                f" da chegada estimada ({chegada_estimada:%H:%M}) considerando o"
                f" deslocamento até {local}: {datas_em_alerta}. As linhas ficam"
                " destacadas em vermelho na grade abaixo — não é um bloqueio,"
                " apenas um alerta."
            )
            if st.checkbox(
                "Já verifiquei e quero manter estes horários assim mesmo",
                key=f"chk_{chave_ignorar}",
            ):
                st.session_state[chave_ignorar] = True
                st.rerun()

    df_exibicao_estilizada = _estilizar_grade_calendario(df_exibicao, horario_minimo_min)

    configuracao = {
        "Data": st.column_config.TextColumn("Data", disabled=True),
        "Dia da Semana": st.column_config.TextColumn("Dia", disabled=True),
        "Status": st.column_config.SelectboxColumn(
            "Status", options=OPCOES_STATUS, required=True
        ),
        "Exam. M": st.column_config.NumberColumn(
            "Exam. M", min_value=0, max_value=50, step=1, default=0
        ),
        "Exam. T": st.column_config.NumberColumn(
            "Exam. T", min_value=0, max_value=50, step=1, default=0
        ),
        "Horário": st.column_config.TextColumn("Horário", disabled=True),
        "Total": st.column_config.NumberColumn("Total", disabled=True),
    }
    for coluna in COLS_VAGAS:
        configuracao[coluna] = st.column_config.NumberColumn(
            coluna, min_value=0, step=1, default=0
        )

    chave_editor = f"editor_{chave}"
    chave_assinatura = f"{chave_editor}__filtro"
    assinatura = (
        tuple(sorted(datas_selecionadas)) if datas_selecionadas else str(periodo)
    )
    if st.session_state.get(chave_assinatura) != assinatura:
        st.session_state.pop(chave_editor, None)
        st.session_state[chave_assinatura] = assinatura

    df_editado = st.data_editor(
        df_exibicao_estilizada,
        column_config=configuracao,
        **LARGURA_TOTAL,
        num_rows="fixed",
        height=450,
        key=chave_editor,
    )
    st.caption(
        "🟩 Linha já lançada (com vagas) · 🟥 Lançada antes do horário sugerido"
        " pelo deslocamento · ⬜ Bloqueada (feriado/indisponível) · sem cor ="
        " ainda vazia."
    )

    # ``update`` ignora NaN: sem o fillna, apagar uma célula descartava a edição.
    colunas_numericas = ["Exam. M", "Exam. T"] + COLS_VAGAS
    df_editado = df_editado.copy()
    for coluna in colunas_numericas:
        df_editado[coluna] = (
            pd.to_numeric(df_editado[coluna], errors="coerce").fillna(0).astype(int)
        )
    df_editado["Status"] = df_editado["Status"].fillna(STATUS_DISPONIVEL)

    df_completo.update(df_editado)
    for coluna in colunas_numericas:
        df_completo[coluna] = (
            pd.to_numeric(df_completo[coluna], errors="coerce").fillna(0).astype(int)
        )

    def _total_da_linha(linha: pd.Series) -> int:
        if linha["Status"] in STATUS_BLOQUEADOS:
            return 0
        return int(sum(int(linha[c]) for c in COLS_VAGAS))

    df_completo["Total"] = df_completo.apply(_total_da_linha, axis=1)
    bloqueadas = df_completo["Status"].isin(STATUS_BLOQUEADOS)
    df_completo.loc[bloqueadas, ["Exam. M", "Exam. T"]] = 0

    st.session_state["historico_localidades"][chave] = df_completo
    salvar_historico(chave, df_completo)

    _monitor_efetivo(banca, mes_nome, ano, datas_do_mes, int(efetivo_maximo))

    st.markdown("---")
    pico = _pico_do_mes(banca, mes_nome, ano, mes_numero)
    col_a, col_b, col_c = st.columns(3)
    col_a.metric("Total Vagas Localidade", f"{int(df_completo['Total'].sum())} vagas")
    col_b.metric("Maior Pico Diário no Mês", f"{pico} / {int(efetivo_maximo)} exam.")
    col_c.metric("Menor Folga Diária", f"{int(efetivo_maximo) - pico} exam.")

    if pico > efetivo_maximo:
        st.error(
            "🚨 **BLOQUEIO DE SEGURANÇA**: o pico de examinadores escalados excede"
            " o limite diário disponível. Ajuste a escala antes de publicar."
        )
        return

    st.markdown("### 🖨️ Conferência e Impressão")
    col_pdf_local, col_pdf_banca = st.columns(2)

    with col_pdf_local:
        st.markdown(f"**Grade detalhada — {local}**")
        bloco_pdf(
            chave=f"loc_{chave}",
            rotulo_gerar="📄 Gerar PDF desta Localidade",
            gerador=lambda: gerar_pdf_localidade(
                df_completo, banca, local, mes_nome, ano
            ),
            nome_arquivo=(
                f"Calendario_{banca}_{local}_{mes_nome}_{ano}.pdf".replace(" ", "_")
            ),
            ajuda="Linha a linha, por horário, apenas desta localidade.",
        )

    with col_pdf_banca:
        st.markdown(f"**Calendário completo — {banca}**")
        dados_banca = dados_da_banca_no_mes(banca, mes_nome, ano)
        st.caption(
            f"{len(dados_banca)} localidade(s) lançada(s) no mês. Ordem do"
            " documento: unidades da sede, região metropolitana e demais"
            " municípios por data do primeiro exame."
        )
        bloco_pdf(
            chave=f"banca_{banca}_{mes_nome}_{ano}",
            rotulo_gerar="📘 Gerar PDF Completo da Banca",
            gerador=lambda: gerar_pdf_banca(
                dados_da_banca_no_mes(banca, mes_nome, ano),
                st.session_state["viagens_registradas"],
                st.session_state["bancas_config"],
                banca,
                mes_nome,
                ano,
                mes_numero,
                int(efetivo_maximo),
            ),
            nome_arquivo=(
                f"Calendario_Completo_{banca}_{mes_nome}_{ano}.pdf".replace(" ", "_")
            ),
            ajuda=(
                "Mostra como o calendário da banca inteira está se desenhando,"
                " com a conferência de efetivo diário."
            ),
        )


# --- Aba: gestão de viagens itinerantes -------------------------------------
def _validar_capacidade(
    bancas_alvo: list[str],
    inicio: datetime.date,
    fim: datetime.date,
    turno: str,
    turnos_por_data: dict[str, str],
    numero_equipe: int,
    examinadores_por_banca: dict[str, int],
    ignorar_id=None,
) -> list[str]:
    """Verifica se o deslocamento simultâneo cabe no limite fora da sede."""
    erros: list[str] = []
    viagens = st.session_state["viagens_registradas"]
    limites = st.session_state["limite_fora_sede_bancas"]
    capacidades = st.session_state["capacidade_bancas"]
    disponibilidade = st.session_state["disponibilidade_semanal"]

    candidata = {
        "Data Inicio": inicio,
        "Data Fim": fim,
        "Turno": turno,
        "Turnos Por Data": turnos_por_data,
    }

    for data_cursor in dias_do_intervalo(inicio, fim, apenas_uteis=True):
        for banca in bancas_alvo:
            quantidade_nova = int(examinadores_por_banca.get(banca, 0))
            picos = []
            for periodo in (TURNO_MANHA, TURNO_TARDE):
                por_equipe: dict[int, int] = {}
                for viagem in viagens:
                    if ignorar_id is not None and viagem.get("id") == ignorar_id:
                        continue
                    if banca_vinculada(viagem, banca) and viagem_cobre_periodo(
                        viagem, data_cursor, periodo
                    ):
                        equipe = int(viagem.get("Numero Banca Itinerante", 0) or 0)
                        por_equipe[equipe] = max(
                            por_equipe.get(equipe, 0),
                            examinadores_da_banca_na_viagem(viagem, banca),
                        )
                if viagem_cobre_periodo(candidata, data_cursor, periodo):
                    por_equipe[int(numero_equipe)] = max(
                        por_equipe.get(int(numero_equipe), 0), quantidade_nova
                    )
                picos.append(sum(por_equipe.values()))

            pico = max(picos or [0])
            limite = int(limites.get(banca, capacidades.get(banca, 0)))
            chave_semana = chave_disponibilidade(banca, data_cursor)
            disponivel = int(disponibilidade.get(chave_semana, capacidades.get(banca, 0)))
            teto = min(limite, disponivel)
            if pico > teto:
                erros.append(
                    f"{banca}: deslocamento simultâneo de {pico} examinadores excede"
                    f" o teto de {teto} em {formatar_br(data_cursor)}."
                )
                return erros
    return erros


def _alertas_de_feriado(
    banca: str, destino: str, inicio: datetime.date, fim: datetime.date
) -> list[datetime.date]:
    texto = st.session_state["feriados_locais_dict"].get(chave_local(banca, destino), "")
    datas = datas_de_feriado_do_texto(texto, inicio.year)
    datas |= datas_de_feriado_do_texto(texto, fim.year)
    return sorted(d for d in datas if inicio <= d <= fim)


def tempo_deslocamento_efetivo(banca: str, local: str) -> int:
    """Minutos de deslocamento a usar: ajuste do usuário, senão a estimativa padrão."""
    chave = chave_local(banca, local)
    valor = st.session_state.get("tempo_deslocamento_dict", {}).get(chave)
    return int(valor) if valor is not None else tempo_deslocamento_padrao(banca, local)


def _painel_sugestao_deslocamento(
    banca: str,
    destino: str,
    horarios_ativos: list[str],
    trechos_da_equipe: list[Viagem],
    data_inicio: datetime.date,
    data_fim: datetime.date,
    turno: str,
) -> None:
    """Mostra o horário sugerido de início considerando o deslocamento.

    Cobre os dois cenários do Ajuste 3: (1) uma cidade só, deslocamento a
    partir da sede da banca; (2) segunda cidade da mesma viagem no mesmo dia,
    descontando 1h de almoço mais o deslocamento entre as duas cidades. Em
    ambos os casos é só uma SUGESTÃO — nada aqui bloqueia o cadastro.
    """
    minutos_base = tempo_deslocamento_efetivo(banca, destino)
    chave_desloc = chave_local(banca, destino)

    if minutos_base <= 0:
        st.caption(
            f"📍 {destino}: sem deslocamento cadastrado (sede/região"
            " metropolitana, ou sem estimativa — informe abaixo se precisar)."
        )

    with st.expander(f"🕒 Deslocamento até {destino}", expanded=minutos_base > 0):
        novo_minutos = st.number_input(
            "Minutos de deslocamento a partir da sede da banca (ajustável):",
            min_value=0,
            max_value=1440,
            value=int(minutos_base),
            step=5,
            key=f"v_desloc_{chave_desloc}",
            help=(
                "Estimativa de referência — corrija aqui se souber o tempo real"
                " de viagem; a correção fica salva para a próxima vez."
            ),
        )
        if novo_minutos != minutos_base:
            st.session_state.setdefault("tempo_deslocamento_dict", {})[
                chave_desloc
            ] = int(novo_minutos)
            salvar_tempo_deslocamento(st.session_state["tempo_deslocamento_dict"])

        if novo_minutos > 0:
            chegada = horario_chegada_estimado(novo_minutos)
            sugestao = sugerir_horario_inicio(horarios_ativos, chegada)
            horas, minutos_resto = divmod(int(novo_minutos), 60)
            duracao_txt = (
                f"{horas}h{minutos_resto:02d}" if horas else f"{minutos_resto}min"
            )
            if sugestao:
                st.info(
                    f"Saindo às {HORA_PARTIDA_PADRAO:%H:%M} da sede, deslocamento de"
                    f" ~{duracao_txt}, chegada estimada às {chegada:%H:%M}."
                    f" **Sugestão: iniciar a primeira turma às {sugestao}.**"
                )
            else:
                st.warning(
                    f"Chegada estimada às {chegada:%H:%M} — depois do último"
                    " horário cadastrado na grade. Considere iniciar o"
                    " atendimento apenas no dia seguinte."
                )

    # Segunda cidade da mesma viagem, no mesmo dia — o caso de Pinheiro pela
    # manhã e São Bento à tarde.
    if data_fim < data_inicio:
        return
    dias_novos = set(dias_do_intervalo(data_inicio, data_fim))
    for trecho in trechos_da_equipe:
        if trecho.get("Destino") == destino:
            continue
        inicio_existente = para_data(trecho.get("Data Inicio"))
        fim_existente = para_data(trecho.get("Data Fim"))
        if not inicio_existente or not fim_existente:
            continue
        dias_comuns = sorted(
            dias_novos & set(dias_do_intervalo(inicio_existente, fim_existente))
        )
        if not dias_comuns:
            continue

        dia_referencia = dias_comuns[0]
        turno_existente = turno_da_viagem_no_dia(trecho, dia_referencia)
        periodo_cidade1 = (
            turno_existente if turno_existente in (TURNO_MANHA, TURNO_TARDE) else TURNO_MANHA
        )
        fim_turno1 = horario_fim_turno(horarios_ativos, periodo_cidade1)
        if not fim_turno1:
            continue

        minutos_destino_existente = tempo_deslocamento_efetivo(
            banca, trecho["Destino"]
        )
        chave_entre = f"v_desloc_entre_{chave_local(banca, destino)}_{chave_local(banca, trecho['Destino'])}"
        deslocamento_padrao_entre = abs(minutos_base - minutos_destino_existente)

        with st.expander(
            f"🕒 Sequência no mesmo dia: {trecho['Destino']} → {destino}",
            expanded=True,
        ):
            if turno_existente == TURNO_INTEGRAL:
                st.caption(
                    f"⚠️ {trecho['Destino']} ainda está marcado como"
                    f" '{TURNO_INTEGRAL}' em {formatar_curto(dia_referencia)}."
                    f" Para atender as duas cidades no mesmo dia, ajuste"
                    f" {trecho['Destino']} para '{TURNO_MANHA}' no ajuste de"
                    " turno por dia (assumindo manhã aqui para calcular)."
                )
            minutos_entre = st.number_input(
                f"Deslocamento estimado entre {trecho['Destino']} e {destino}"
                " (minutos):",
                min_value=0,
                max_value=600,
                value=int(deslocamento_padrao_entre),
                step=5,
                key=chave_entre,
                help=(
                    "Sem uma rota exata entre as duas cidades, o padrão é a"
                    " diferença entre os tempos de cada uma até a sede — ajuste"
                    " se souber a distância real entre elas."
                ),
            )
            sugestao2, chegada2 = sugerir_horario_segunda_cidade(
                horarios_ativos, fim_turno1, minutos_entre
            )
            if sugestao2:
                st.info(
                    f"{trecho['Destino']} termina ~{fim_turno1:%H:%M} · +1h de"
                    f" almoço · +{minutos_entre}min até {destino} → chegada"
                    f" estimada às {chegada2:%H:%M}. **Sugestão: iniciar"
                    f" {destino} às {sugestao2}"
                    f" ({TURNO_TARDE if periodo_cidade1 == TURNO_MANHA else TURNO_MANHA}).**"
                )
            else:
                st.warning(
                    f"Chegada estimada às {chegada2:%H:%M} — depois do último"
                    f" horário e do limite de {LIMITE_FIM_EXPEDIENTE:%H:%M}. Talvez"
                    " não dê para atender as duas cidades no mesmo dia."
                )
        break  # um só bloco de sugestão, com o primeiro trecho que colide


def _editor_turnos_por_dia(
    inicio: datetime.date,
    fim: datetime.date,
    turno_geral: str,
    prefixo_chave: str,
    valores_atuais: dict[str, str] | None = None,
    expandido: bool = False,
) -> dict[str, str]:
    """Permite definir o turno dia a dia dentro do período da viagem.

    É o que resolve o caso de a equipe atender Pinheiro pela manhã e São Bento
    à tarde no dia 12: sem isso, o turno valia para todos os dias do trecho e
    o quadro contava 6 + 6 examinadores no mesmo dia.
    """
    valores_atuais = valores_atuais or {}
    dias = dias_do_intervalo(inicio, fim)
    if not dias:
        return {}
    if len(dias) > 45:
        st.warning("Período muito longo para o ajuste dia a dia.")
        return dict(valores_atuais)

    # Os seletores guardam estado por chave. Se o período ou o turno geral
    # mudam, os valores antigos ficariam presos e o sistema acusaria conflito
    # onde não há — por isso o estado é descartado quando o contexto muda.
    chave_contexto = f"{prefixo_chave}__contexto"
    contexto = (inicio.isoformat(), fim.isoformat(), turno_geral)
    if st.session_state.get(chave_contexto) != contexto:
        for chave_antiga in [
            k
            for k in st.session_state
            if k.startswith(f"{prefixo_chave}_") and k != chave_contexto
        ]:
            st.session_state.pop(chave_antiga, None)
        st.session_state[chave_contexto] = contexto

    resultado: dict[str, str] = {}
    with st.expander(
        f"🗓️ Ajuste de turno por dia ({len(dias)} dia(s))", expanded=expandido
    ):
        st.caption(
            "Por padrão todos os dias seguem o período geral da viagem. Altere"
            " apenas os dias em que a equipe atende só um turno, ou use"
            f" **{TURNO_SEM_ATENDIMENTO}** para dias de deslocamento e fins de"
            " semana."
        )
        colunas = st.columns(min(len(dias), 3))
        for indice, dia in enumerate(dias):
            iso = dia.isoformat()
            atual = valores_atuais.get(iso, turno_geral)
            if atual not in TURNOS_POR_DIA:
                atual = turno_geral
            with colunas[indice % len(colunas)]:
                escolhido = st.selectbox(
                    f"{formatar_curto(dia)} ({nome_dia_semana(dia)[:3]})",
                    TURNOS_POR_DIA,
                    index=TURNOS_POR_DIA.index(atual),
                    key=f"{prefixo_chave}_{iso}",
                )
            if escolhido != turno_geral:
                resultado[iso] = escolhido
    return resultado


def _formulario_cadastro_viagem() -> None:
    st.markdown("#### ➕ Cadastrar Nova Viagem")

    bancas_config = st.session_state["bancas_config"]
    if not bancas_config:
        st.warning("Cadastre ao menos uma banca na aba de Gestão de Localidades.")
        return

    viagens = st.session_state["viagens_registradas"]
    banca = st.selectbox("Banca Responsável:", list(bancas_config.keys()), key="v_banca")

    numeros_existentes = numeros_de_equipe(viagens, banca)
    proximo = proximo_numero_equipe(viagens, banca)
    opcoes_equipe = {f"➕ Nova banca itinerante (Banca {proximo:02d})": proximo}
    for numero in numeros_existentes:
        opcoes_equipe[f"🔁 Continuar Banca {numero:02d} já cadastrada"] = numero

    escolha = st.selectbox(
        "Esta viagem pertence a:",
        list(opcoes_equipe.keys()),
        key="v_equipe",
        help=(
            "Se a mesma equipe vai atender mais de um município na mesma viagem,"
            " cadastre o primeiro trecho como 'Nova banca itinerante' e os"
            " seguintes como 'Continuar' essa mesma banca. Assim o sistema conta"
            " os examinadores uma única vez."
        ),
    )
    numero_equipe = opcoes_equipe[escolha]

    trecho_referencia = next(
        (
            v
            for v in viagens
            if v.get("Banca") == banca
            and int(v.get("Numero Banca Itinerante", 0) or 0) == numero_equipe
        ),
        None,
    )
    padrao_examinadores = int(
        (trecho_referencia or {}).get("Examinadores")
        or st.session_state["limite_fora_sede_bancas"].get(
            banca, st.session_state["capacidade_bancas"].get(banca, 1)
        )
        or 1
    )

    bancas_apoio: list[str] = []
    possiveis = bancas_apoio_disponiveis(banca)
    if possiveis:
        bancas_apoio = st.multiselect(
            "🤝 Bancas em apoio conjunto:",
            possiveis,
            key="v_apoio",
            help="O efetivo de cada banca é informado separadamente abaixo.",
        )

    sede = obter_sede(bancas_config, banca)
    destinos = [l for l in bancas_config[banca] if l != sede] or bancas_config[banca]
    if not destinos:
        st.warning("Esta banca não possui localidades cadastradas.")
        return
    destino = st.selectbox("Município de Destino:", destinos, key="v_destino")

    hoje = datetime.date.today()
    col_ini, col_fim = st.columns(2)
    data_inicio = col_ini.date_input(
        "Início da Viagem:", hoje, format="DD/MM/YYYY", key="v_inicio"
    )
    data_fim = col_fim.date_input(
        "Término da Viagem:",
        hoje + datetime.timedelta(days=4),
        format="DD/MM/YYYY",
        key="v_fim",
    )

    turno = st.selectbox(
        "Período de atendimento (padrão do trecho):",
        TURNOS,
        key="v_turno",
        help=(
            "Este é o padrão aplicado a todos os dias. Para dias em que a equipe"
            " atende só um turno, use o ajuste dia a dia logo abaixo."
        ),
    )

    turnos_por_data: dict[str, str] = {}
    if data_fim >= data_inicio:
        turnos_por_data = _editor_turnos_por_dia(
            data_inicio, data_fim, turno, "v_turno_dia"
        )
        if turnos_por_data:
            st.info(
                "Ajustes aplicados: "
                + " · ".join(
                    f"{formatar_curto(para_data(iso))} → {valor}"
                    for iso, valor in sorted(turnos_por_data.items())
                )
            )

    horarios_ativos = [
        h
        for h in st.session_state["lista_horarios"]
        if h not in st.session_state["horarios_inativos"]
    ]
    trechos_da_equipe = [
        v
        for v in viagens
        if v.get("Banca") == banca
        and int(v.get("Numero Banca Itinerante", 0) or 0) == numero_equipe
    ]
    if horarios_ativos and data_fim >= data_inicio:
        _painel_sugestao_deslocamento(
            banca, destino, horarios_ativos, trechos_da_equipe, data_inicio, data_fim, turno
        )

    examinadores_por_banca: dict[str, int] = {}
    participantes = bancas_para_validar(banca, bancas_apoio)
    if bancas_apoio:
        colunas = st.columns(len(participantes))
        for coluna, participante in zip(colunas, participantes):
            limite = int(st.session_state["limite_fora_sede_bancas"].get(participante, 50))
            examinadores_por_banca[participante] = coluna.number_input(
                f"Examinadores — {participante.replace('Banca ', '')}:",
                min_value=0,
                max_value=max(limite, 1),
                value=min(padrao_examinadores, max(limite, 1)),
                step=1,
                key=f"v_exam_{participante}",
            )
    else:
        examinadores_por_banca[banca] = st.number_input(
            "Nº de Examinadores Deslocados:",
            min_value=1,
            max_value=50,
            value=max(1, min(50, padrao_examinadores)),
            step=1,
            key="v_exam_unico",
        )

    observacoes = st.text_input(
        "Observações / Portaria:", placeholder="Ex: Portaria nº 123/2026", key="v_obs"
    )

    # Apenas ``key``: o valor inicial já veio do banco. Passar ``value`` junto
    # com ``key`` gera warning do Streamlit e conflita ao restaurar backup.
    controle_ativo = st.checkbox(
        "Ativar controle automático de capacidade/disponibilidade",
        key="controle_capacidade_ativo",
        help="Desativado: permite cadastrar viagens sem bloqueio por capacidade.",
    )
    salvar_config("controle_capacidade_ativo", controle_ativo)

    _painel_capacidade()

    if st.button("💾 Registrar Viagem", key="v_btn_salvar"):
        _registrar_viagem(
            banca=banca,
            numero_equipe=numero_equipe,
            bancas_apoio=bancas_apoio,
            destino=destino,
            data_inicio=data_inicio,
            data_fim=data_fim,
            turno=turno,
            turnos_por_data=turnos_por_data,
            examinadores_por_banca=examinadores_por_banca,
            observacoes=observacoes,
            controle_ativo=controle_ativo,
        )


def _painel_capacidade() -> None:
    with st.expander("⚙️ Capacidade e disponibilidade semanal", expanded=False):
        st.caption("Ajuste manualmente os limites operacionais de cada banca.")
        capacidades = st.session_state["capacidade_bancas"]
        limites = st.session_state["limite_fora_sede_bancas"]
        colunas = st.columns(2)

        for indice, banca in enumerate(st.session_state["bancas_config"].keys()):
            with colunas[indice % 2]:
                capacidades[banca] = st.number_input(
                    f"{banca} — efetivo total",
                    min_value=0,
                    max_value=100,
                    value=int(capacidades.get(banca, 0)),
                    key=f"cap_{banca}",
                )
                limites[banca] = st.number_input(
                    f"{banca} — máximo fora da sede",
                    min_value=0,
                    max_value=100,
                    value=int(limites.get(banca, 0)),
                    key=f"lim_{banca}",
                )
        salvar_config("capacidade_bancas", capacidades)
        salvar_config("limite_fora_sede_bancas", limites)

        st.markdown("**Disponibilidade semanal**")
        semana_referencia = st.date_input(
            "Semana de referência (qualquer dia):",
            datetime.date.today(),
            format="DD/MM/YYYY",
            key="v_semana_ref",
        )
        segunda = semana_referencia - datetime.timedelta(days=semana_referencia.weekday())
        disponibilidade = st.session_state["disponibilidade_semanal"]
        bancas_disp = [
            b
            for b in BANCAS_COM_DISPONIBILIDADE_SEMANAL
            if b in st.session_state["bancas_config"]
        ]
        if bancas_disp:
            colunas_disp = st.columns(len(bancas_disp))
            for coluna, banca in zip(colunas_disp, bancas_disp):
                chave = chave_disponibilidade(banca, segunda)
                with coluna:
                    disponibilidade[chave] = st.number_input(
                        f"{banca} — disponível em {formatar_br(segunda)}",
                        min_value=0,
                        max_value=100,
                        value=int(disponibilidade.get(chave, capacidades.get(banca, 0))),
                        key=f"disp_{banca}_{segunda.isoformat()}",
                    )
            salvar_disponibilidade(disponibilidade)


def _registrar_viagem(
    *,
    banca: str,
    numero_equipe: int,
    bancas_apoio: list[str],
    destino: str,
    data_inicio: datetime.date,
    data_fim: datetime.date,
    turno: str,
    turnos_por_data: dict[str, str],
    examinadores_por_banca: dict[str, int],
    observacoes: str,
    controle_ativo: bool,
) -> None:
    erros: list[str] = []

    if data_fim < data_inicio:
        erros.append("A data final não pode ser anterior à data de início.")

    total = sum(int(v or 0) for v in examinadores_por_banca.values())
    if total <= 0:
        erros.append("Informe ao menos um examinador para a viagem.")

    candidata = {
        "Data Inicio": data_inicio,
        "Data Fim": data_fim,
        "Turno": turno,
        "Turnos Por Data": turnos_por_data,
    }
    if data_fim >= data_inicio and not dias_efetivos_da_viagem(candidata):
        erros.append(
            "Todos os dias do período estão marcados como"
            f" '{TURNO_SEM_ATENDIMENTO}'. A viagem não teria atendimento."
        )

    conflitos = conflitos_de_rota(
        st.session_state["viagens_registradas"],
        banca=banca,
        numero_equipe=numero_equipe,
        inicio=data_inicio,
        fim=data_fim,
        turno=turno,
        turnos_por_data=turnos_por_data,
    )
    if conflitos:
        destinos = ", ".join(sorted({str(c.get("Destino", "?")) for c in conflitos}))
        erros.append(
            f"Esta equipe já atende {destinos} no mesmo turno de algum dia deste"
            " período. Use o ajuste de turno por dia para separar manhã e tarde."
        )

    if controle_ativo:
        erros.extend(
            _validar_capacidade(
                bancas_para_validar(banca, bancas_apoio),
                data_inicio,
                data_fim,
                turno,
                turnos_por_data,
                numero_equipe,
                examinadores_por_banca,
            )
        )

    feriados = _alertas_de_feriado(banca, destino, data_inicio, data_fim)

    if erros:
        for erro in erros:
            st.error(erro)
        return

    registro = {
        "id": None,
        "Banca": banca,
        "Numero Banca Itinerante": int(numero_equipe),
        "Bancas Apoio": list(bancas_apoio),
        "Destino": destino,
        "Data Inicio": data_inicio,
        "Data Fim": data_fim,
        "Turno": turno,
        "Turnos Por Data": dict(turnos_por_data),
        "Examinadores": total,
        "Examinadores Por Banca": {k: int(v) for k, v in examinadores_por_banca.items()},
        "Observações": observacoes or "",
    }
    st.session_state["viagens_registradas"].append(registro)
    salvar_viagens(st.session_state["viagens_registradas"])

    if feriados:
        st.warning(
            "⚠️ A viagem atravessa feriado(s) municipal(is): "
            + ", ".join(formatar_br(d) for d in feriados)
        )
    st.success(
        f"Trecho para **{destino}** registrado na Banca {int(numero_equipe):02d}."
    )
    st.rerun()


def _bloco_equipe(banca: str, numero: int, trechos: list[Viagem]) -> None:
    destinos: list[str] = []
    for trecho in trechos:
        destino = trecho.get("Destino", "")
        if destino and destino not in destinos:
            destinos.append(destino)
    nome_destinos = " e ".join(destinos) if destinos else "Destino não informado"

    referencia = trechos[0]
    st.markdown(f"**📍 {nome_destinos}** *(Banca Itinerante {numero:02d})*")

    col_ini, col_fim, col_num, col_exam = st.columns([2, 2, 1.5, 1.5])
    sufixo = f"{banca}_{numero}"

    novo_inicio = col_ini.date_input(
        "Início",
        value=para_data(referencia["Data Inicio"]),
        format="DD/MM/YYYY",
        key=f"grp_ini_{sufixo}",
    )
    novo_fim = col_fim.date_input(
        "Término",
        value=para_data(referencia["Data Fim"]),
        format="DD/MM/YYYY",
        key=f"grp_fim_{sufixo}",
    )
    novo_numero = col_num.number_input(
        "Nº Banca", value=int(numero), min_value=1, step=1, key=f"grp_num_{sufixo}"
    )
    novo_efetivo = col_exam.number_input(
        "Examinadores",
        value=int(referencia.get("Examinadores", 1) or 1),
        min_value=1,
        step=1,
        key=f"grp_exam_{sufixo}",
    )

    mudou = (
        novo_inicio != para_data(referencia["Data Inicio"])
        or novo_fim != para_data(referencia["Data Fim"])
        or int(novo_numero) != int(numero)
        or int(novo_efetivo) != int(referencia.get("Examinadores", 1) or 1)
    )
    if mudou:
        if novo_fim < novo_inicio:
            st.error("A data final não pode ser anterior à data de início.")
        else:
            for trecho in trechos:
                trecho["Data Inicio"] = novo_inicio
                trecho["Data Fim"] = novo_fim
                trecho["Numero Banca Itinerante"] = int(novo_numero)
                trecho["Examinadores"] = int(novo_efetivo)
                trecho["Turnos Por Data"] = {
                    iso: valor
                    for iso, valor in (trecho.get("Turnos Por Data") or {}).items()
                    if (d := para_data(iso)) and novo_inicio <= d <= novo_fim
                }
                por_banca = trecho.get("Examinadores Por Banca") or {}
                if len(por_banca) == 1:
                    unica = next(iter(por_banca))
                    trecho["Examinadores Por Banca"] = {unica: int(novo_efetivo)}
            salvar_viagens(st.session_state["viagens_registradas"])
            st.rerun()

    for indice, trecho in enumerate(trechos):
        ajustes = descricao_turnos(trecho)
        rotulo = (
            f"🛠️ {trecho.get('Destino')} — {trecho.get('Turno', TURNO_INTEGRAL)}"
            + (f" · ajustes: {ajustes}" if ajustes else "")
        )
        with st.expander(rotulo, expanded=False):
            inicio_trecho = para_data(trecho["Data Inicio"])
            fim_trecho = para_data(trecho["Data Fim"])
            novo_turno = st.selectbox(
                "Período padrão deste trecho:",
                TURNOS,
                index=TURNOS.index(trecho.get("Turno", TURNO_INTEGRAL))
                if trecho.get("Turno", TURNO_INTEGRAL) in TURNOS
                else 0,
                key=f"trecho_turno_{sufixo}_{indice}",
            )
            novos_ajustes = _editor_turnos_por_dia(
                inicio_trecho,
                fim_trecho,
                novo_turno,
                f"trecho_dia_{sufixo}_{indice}",
                trecho.get("Turnos Por Data"),
                expandido=True,
            )
            if novo_turno != trecho.get("Turno", TURNO_INTEGRAL) or novos_ajustes != (
                trecho.get("Turnos Por Data") or {}
            ):
                if st.button(
                    "💾 Aplicar turnos deste trecho",
                    key=f"trecho_salvar_{sufixo}_{indice}",
                ):
                    trecho["Turno"] = novo_turno
                    trecho["Turnos Por Data"] = novos_ajustes
                    salvar_viagens(st.session_state["viagens_registradas"])
                    st.rerun()
            if trecho.get("Observações"):
                st.caption(str(trecho["Observações"]))

    col_remover, _ = st.columns([3, 7])
    if col_remover.button(f"🗑️ Remover Banca {numero:02d}", key=f"del_{sufixo}"):
        salvar_snapshot(motivo="antes_remover_viagem")
        st.session_state["viagens_registradas"] = [
            v
            for v in st.session_state["viagens_registradas"]
            if not (
                v.get("Banca") == banca
                and int(v.get("Numero Banca Itinerante", 1) or 1) == int(numero)
            )
        ]
        salvar_viagens(st.session_state["viagens_registradas"])
        st.rerun()
    st.markdown("---")


def _cronograma_viagens() -> None:
    st.markdown("### 📋 Cronograma de Viagens Programadas")
    viagens = st.session_state["viagens_registradas"]
    if not viagens:
        st.warning("Nenhuma viagem cadastrada até o momento.")
        return

    col_mes, col_ano = st.columns(2)
    mes_exibicao = col_mes.selectbox(
        "Mês de referência:", MESES_LISTA, index=indice_mes_atual(), key="v_mes_cronograma"
    )
    ano_exibicao = col_ano.selectbox(
        "Ano de referência:", anos_disponiveis(), index=0, key="v_ano_cronograma"
    )
    mes_numero = numero_do_mes(mes_exibicao)

    st.markdown("#### 🖨️ Escala mensal de viagens")
    st.caption(
        "O PDF atrela automaticamente 01 preposto e 01 veículo a cada banca"
        " itinerante — essa informação existe apenas no documento."
    )
    bloco_pdf(
        chave=f"escala_viagens_{mes_exibicao}_{ano_exibicao}",
        rotulo_gerar="🚍 Gerar PDF da Escala de Viagens do Mês",
        gerador=lambda: gerar_pdf_escala_viagens(
            st.session_state["viagens_registradas"],
            mes_exibicao,
            ano_exibicao,
            mes_numero,
        ),
        nome_arquivo=f"Escala_Viagens_{mes_exibicao}_{ano_exibicao}.pdf",
        ajuda="Consolida todas as equipes itinerantes do mês.",
        altura_previa=560,
    )
    st.markdown("---")

    for banca in sorted({v.get("Banca", "") for v in viagens if v.get("Banca")}):
        st.markdown(
            f'<div class="faixa-banca">MÊS DE {html.escape(mes_exibicao).upper()}'
            f" — {html.escape(banca).upper()}</div>",
            unsafe_allow_html=True,
        )
        viagens_da_banca = [v for v in viagens if v.get("Banca") == banca]
        for numero in sorted(
            {int(v.get("Numero Banca Itinerante", 1) or 1) for v in viagens_da_banca}
        ):
            trechos = [
                v
                for v in viagens_da_banca
                if int(v.get("Numero Banca Itinerante", 1) or 1) == numero
            ]
            trechos.sort(key=lambda v: para_data(v["Data Inicio"]))
            _bloco_equipe(banca, numero, trechos)


def aba_viagens() -> None:
    st.markdown("### 🚍 Controle de Viagens e Equipes Itinerantes")
    st.info(
        "Organize aqui os deslocamentos das bancas. As equipes em viagem são"
        " somadas ao controle de efetivo diário e aplicadas na montagem do"
        " calendário. Cada equipe é contada uma única vez por dia, mesmo"
        " atendendo dois municípios em turnos diferentes."
    )
    coluna_form, coluna_lista = st.columns([1, 2])
    with coluna_form:
        _formulario_cadastro_viagem()
    with coluna_lista:
        _cronograma_viagens()


# --- Aba: horários e turmas -------------------------------------------------
PADRAO_HORARIO = re.compile(r"^([01]?\d|2[0-3]):([0-5]\d)$")


def _normalizar_horario(texto: str) -> str | None:
    bruto = (texto or "").strip().replace("h", ":").replace(".", ":")
    correspondencia = PADRAO_HORARIO.match(bruto)
    if not correspondencia:
        return None
    hora, minuto = correspondencia.groups()
    return f"{int(hora):02d}:{minuto}"


def _persistir_horarios(lista: list[str], inativos: list[str]) -> None:
    salvar_config("lista_horarios", lista)
    salvar_config("horarios_inativos", inativos)


def aba_horarios() -> None:
    st.markdown("### ⏰ Grade de Horários e Turmas")
    st.info(
        "Os horários definidos aqui são as linhas geradas para cada dia de"
        " atendimento na aba **Montar Calendário Detalhado**. Horários"
        " desativados param de ser oferecidos em novos calendários, mas os"
        " lançamentos já feitos são preservados."
    )

    lista: list[str] = st.session_state["lista_horarios"]
    inativos: list[str] = st.session_state["horarios_inativos"]
    col_cadastro, col_visao = st.columns([1, 2])

    with col_cadastro:
        st.markdown("#### ➕ Adicionar Horário")
        novo = st.text_input("Horário (HH:MM):", placeholder="Ex: 17:30", key="hora_novo")
        if st.button("➕ Adicionar", key="hora_btn_add"):
            normalizado = _normalizar_horario(novo)
            if not normalizado:
                st.error("Formato inválido. Use HH:MM, por exemplo 14:30.")
            elif normalizado in lista:
                st.warning("Este horário já está cadastrado.")
            else:
                lista.append(normalizado)
                lista.sort()
                _persistir_horarios(lista, inativos)
                st.success(f"Horário **{normalizado}** adicionado.")
                st.rerun()

        st.markdown("---")
        st.markdown("#### ♻️ Restaurar Grade Padrão")
        st.caption("Volta aos oito horários originais. Não altera calendários já lançados.")
        if st.button("♻️ Restaurar padrão", key="hora_btn_padrao"):
            st.session_state["lista_horarios"] = list(HORARIOS_PADRAO)
            st.session_state["horarios_inativos"] = []
            _persistir_horarios(
                st.session_state["lista_horarios"], st.session_state["horarios_inativos"]
            )
            st.success("Grade padrão restaurada.")
            st.rerun()

    with col_visao:
        st.markdown("#### 🗂️ Horários Cadastrados")
        if not lista:
            st.warning("Nenhum horário cadastrado. Os calendários sairiam vazios.")
            return

        ativos = [h for h in lista if h not in inativos]
        df_visao = pd.DataFrame(
            {
                "Horário": lista,
                "Período": [periodo_do_horario(h) for h in lista],
                "Ativo": [h not in inativos for h in lista],
            }
        ).sort_values("Horário", ignore_index=True)

        editado = st.data_editor(
            df_visao,
            hide_index=True,
            **LARGURA_TOTAL,
            num_rows="fixed",
            column_config={
                "Horário": st.column_config.TextColumn("Horário", disabled=True),
                "Período": st.column_config.TextColumn("Período", disabled=True),
                "Ativo": st.column_config.CheckboxColumn(
                    "Ativo", help="Desmarque para parar de oferecer este horário."
                ),
            },
            key="hora_editor",
        )

        novos_inativos = sorted(
            editado.loc[~editado["Ativo"].fillna(True), "Horário"].tolist()
        )
        if novos_inativos != sorted(inativos):
            st.session_state["horarios_inativos"] = novos_inativos
            _persistir_horarios(lista, novos_inativos)
            st.rerun()

        st.markdown("---")
        col_m, col_t, col_total = st.columns(3)
        col_m.metric(
            "Turmas de Manhã",
            sum(1 for h in ativos if periodo_do_horario(h) == TURNO_MANHA),
        )
        col_t.metric(
            "Turmas de Tarde",
            sum(1 for h in ativos if periodo_do_horario(h) == TURNO_TARDE),
        )
        col_total.metric("Horários ativos", f"{len(ativos)} / {len(lista)}")

        st.markdown("#### 🗑️ Excluir Horário Definitivamente")
        alvo = st.selectbox("Horário:", lista, key="hora_excluir")
        st.caption("Para apenas suspender o horário, use a coluna *Ativo* acima.")
        if st.button("🗑️ Excluir", key="hora_btn_del"):
            lista.remove(alvo)
            if alvo in inativos:
                inativos.remove(alvo)
            _persistir_horarios(lista, inativos)
            st.warning(f"Horário **{alvo}** excluído da grade.")
            st.rerun()


# --- Aba: dashboard consolidado ---------------------------------------------
def _linha_resumo(
    banca: str, local: str, mes: str, ano: str, df_local: pd.DataFrame
) -> dict | None:
    df_vagas = df_local[
        (df_local["Status"] == STATUS_DISPONIVEL) & (df_local["Total"] > 0)
    ]
    if df_vagas.empty:
        return None

    por_data = df_vagas.groupby("Data")[["Exam. M", "Exam. T"]].max()
    pico = int(
        por_data.apply(
            lambda linha: pico_diario(_num(linha["Exam. M"]), _num(linha["Exam. T"])),
            axis=1,
        ).max()
    )
    dias = [d.day for d in (para_data(x) for x in df_vagas["Data"].unique()) if d]

    resumo = {
        "Banca": banca,
        "Localidade": local,
        "Mês/Ano": f"{mes}/{ano}",
        "Dias com Exame": int(df_vagas["Data"].nunique()),
        "Datas dos exames": formatar_datas_exames(dias),
        "Pico Exam./Dia": pico,
        "Total Vagas": int(df_vagas["Total"].fillna(0).sum()),
    }
    for coluna in COLS_VAGAS:
        resumo[coluna] = (
            int(df_vagas[coluna].fillna(0).sum()) if coluna in df_vagas.columns else 0
        )
    return resumo


def aba_relatorio() -> None:
    st.markdown("### 📊 Dashboard Executivo de Oferta de Vagas")

    bancas_config = st.session_state["bancas_config"]
    col_mes, col_ano, col_banca, col_local = st.columns(4)
    mes_filtro = col_mes.selectbox(
        "🗓️ Mês:", MESES_LISTA, index=indice_mes_atual(), key="rel_mes"
    )
    ano_filtro = col_ano.selectbox("📅 Ano:", anos_disponiveis(), index=0, key="rel_ano")
    banca_filtro = col_banca.selectbox(
        "🏛️ Banca:", ["Todas"] + list(bancas_config.keys()), key="rel_banca"
    )
    opcoes_local = (
        bancas_config.get(banca_filtro, [])
        if banca_filtro != "Todas"
        else todas_localidades(bancas_config)
    )
    local_filtro = col_local.selectbox(
        "📍 Localidade:", ["Todas as Localidades"] + opcoes_local, key="rel_local"
    )

    st.markdown("---")

    resumos = []
    for chave, df_local in st.session_state["historico_localidades"].items():
        partes = desmontar_chave_historico(chave)
        if not partes:
            LOG.warning("Chave de histórico ignorada por formato inválido: %s", chave)
            continue
        banca, mes, ano, local = partes
        if mes != mes_filtro or ano != str(ano_filtro):
            continue
        if banca_filtro != "Todas" and banca != banca_filtro:
            continue
        if local_filtro != "Todas as Localidades" and local != local_filtro:
            continue
        if df_local is None or df_local.empty:
            continue
        resumo = _linha_resumo(banca, local, mes, ano, df_local)
        if resumo:
            resumos.append(resumo)

    if not resumos:
        st.info("Nenhum dado cadastrado ou encontrado para os filtros selecionados.")
        return

    df_resumo = pd.DataFrame(resumos).sort_values(
        ["Banca", "Localidade"], ignore_index=True
    )

    st.markdown("#### 📈 Resumo Geral da Seleção")
    total_vagas = int(df_resumo["Total Vagas"].sum())
    total_dias = int(df_resumo["Dias com Exame"].sum())
    media = round(total_vagas / total_dias, 1) if total_dias else 0
    pico_maximo = int(df_resumo["Pico Exam./Dia"].max())

    kpi1, kpi2, kpi3, kpi4 = st.columns(4)
    kpi1.metric("🎯 Total Vagas Ofertadas", f"{total_vagas} vagas")
    kpi2.metric("📅 Total Dias com Exame", f"{total_dias} dias")
    kpi3.metric("📊 Média Vagas / Dia", f"{media} vagas")
    kpi4.metric("👨‍⚖️ Maior Pico Examinadores", f"{pico_maximo} exam.")

    total_categorias = int(df_resumo[COLS_CATEGORIA].sum().sum())
    total_pcd = int(df_resumo[COLS_PCD].sum().sum())
    if total_pcd:
        st.caption(
            f"Composição: {total_categorias} vagas de categoria comum e"
            f" {total_pcd} vagas PCD."
        )

    st.markdown("---")
    st.markdown("#### 🏢 Panorama Detalhado por Localidade")
    st.dataframe(df_resumo, **LARGURA_TOTAL, hide_index=True)

    incluir_pcd = st.checkbox(
        "Incluir colunas PCD na planilha de publicação",
        value=bool(total_pcd),
        key="rel_incluir_pcd",
        help=(
            "O modelo oficial antigo trazia apenas as categorias A–E. Mantenha"
            " marcado para que os totais da planilha batam com os do painel."
        ),
    )
    excel = gerar_excel_publicacao(df_resumo, mes_filtro, ano_filtro, incluir_pcd)
    st.download_button(
        label="📥 Exportar Relatório Consolidado (Excel Oficial)",
        data=excel,
        file_name=f"Calendario_Publicacao_{mes_filtro.upper()}_{ano_filtro}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="rel_download_excel",
    )


# --- Aba: gestão de bancas e localidades ------------------------------------
def _limpar_dados_da_localidade(banca: str, local: str) -> None:
    """Remove tudo que está pendurado em uma localidade."""
    remover_historico_do_local(banca, local)

    chave = chave_local(banca, local)
    st.session_state["dias_permitidos_dict"].pop(chave, None)
    st.session_state["feriados_locais_dict"].pop(chave, None)

    restantes = [
        v
        for v in st.session_state["viagens_registradas"]
        if not (banca_vinculada(v, banca) and v.get("Destino") == local)
    ]
    removidas = len(st.session_state["viagens_registradas"]) - len(restantes)
    st.session_state["viagens_registradas"] = restantes

    salvar_dias_permitidos(st.session_state["dias_permitidos_dict"])
    salvar_feriados_locais(st.session_state["feriados_locais_dict"])
    if removidas:
        salvar_viagens(st.session_state["viagens_registradas"])
    LOG.info(
        "Localidade %s/%s removida (%d viagens excluídas)", banca, local, removidas
    )


def _contar_dados_da_banca(banca: str) -> tuple[int, int]:
    calendarios = sum(
        1
        for chave in st.session_state["historico_localidades"]
        if (partes := desmontar_chave_historico(chave)) and partes[0] == banca
    )
    viagens = sum(
        1 for v in st.session_state["viagens_registradas"] if v.get("Banca") == banca
    )
    return calendarios, viagens


def aba_gestao() -> None:
    st.markdown("### ➕ Gestão de Bancas e Localidades")
    st.info(
        "Adicione ou remova Bancas e Municípios. A remoção apaga também os"
        " calendários, feriados e viagens vinculados."
    )

    bancas_config: dict[str, list[str]] = st.session_state["bancas_config"]
    col_banca, col_local = st.columns(2)

    with col_banca:
        st.markdown("#### 🏛️ Cadastrar Nova Banca")
        nova_banca = st.text_input(
            "Nome da Nova Banca:", placeholder="Ex: Banca Balsas", key="gestao_nova_banca"
        )
        if st.button("➕ Adicionar Banca", key="gestao_btn_add_banca"):
            nome = (nova_banca or "").strip()
            if not nome:
                st.warning("Informe um nome para a banca.")
            elif nome in bancas_config:
                st.warning("Esta banca já existe no sistema.")
            else:
                bancas_config[nome] = []
                salvar_bancas(bancas_config)
                st.success(f"Banca **{nome}** criada com sucesso!")
                st.rerun()

    with col_local:
        st.markdown("#### 📍 Cadastrar Novo Município / Localidade")
        if not bancas_config:
            st.warning("Cadastre uma banca antes de adicionar localidades.")
        else:
            banca_destino = st.selectbox(
                "Selecione a Banca que receberá o local:",
                list(bancas_config.keys()),
                key="gestao_banca_destino",
            )
            novo_local = st.text_input(
                "Nome do Município/Localidade:",
                placeholder="Ex: Balsas - Pátio",
                key="gestao_novo_local",
            )
            st.caption(
                "A primeira localidade da banca é tratada como sede e aparece"
                " primeiro nos relatórios em PDF."
            )
            if st.button("➕ Adicionar Localidade", key="gestao_btn_add_local"):
                nome = (novo_local or "").strip()
                if not nome:
                    st.warning("Informe um nome para a localidade.")
                elif nome in bancas_config[banca_destino]:
                    st.warning("Esta localidade já está cadastrada nesta banca.")
                else:
                    bancas_config[banca_destino].append(nome)
                    salvar_bancas(bancas_config)
                    st.success(f"Localidade **{nome}** vinculada à **{banca_destino}**!")
                    st.rerun()

    st.markdown("---")
    st.markdown("#### 🗑️ Remover Localidade Existente")
    col_banca_rem, col_local_rem = st.columns(2)

    if not bancas_config:
        st.info("Nenhuma banca cadastrada.")
        return

    with col_banca_rem:
        banca_remocao = st.selectbox(
            "Selecione a Banca:", list(bancas_config.keys()), key="gestao_banca_remocao"
        )

    with col_local_rem:
        locais = bancas_config.get(banca_remocao, [])
        if not locais:
            st.info(f"A banca **{banca_remocao}** não possui localidades cadastradas.")
        else:
            local_remocao = st.selectbox(
                "Selecione a Localidade para Remover:",
                locais,
                key="gestao_local_remocao",
            )
            confirmar = st.checkbox(
                f"Confirmo a exclusão de **{local_remocao}** e de todos os dados"
                " vinculados a ela.",
                key="gestao_confirma_local",
            )
            if st.button("🗑️ Remover Localidade", key="gestao_btn_rem_local"):
                if not confirmar:
                    st.warning("Marque a confirmação antes de remover.")
                else:
                    salvar_snapshot(motivo="antes_remover_localidade")
                    _limpar_dados_da_localidade(banca_remocao, local_remocao)
                    bancas_config[banca_remocao].remove(local_remocao)
                    salvar_bancas(bancas_config)
                    st.warning(
                        f"Localidade **{local_remocao}** e seus dados foram removidos."
                    )
                    st.rerun()

    st.markdown("---")
    with st.expander("⚠️ Remover uma Banca inteira", expanded=False):
        banca_excluir = st.selectbox(
            "Banca a excluir:", list(bancas_config.keys()), key="gestao_banca_excluir"
        )
        calendarios, viagens = _contar_dados_da_banca(banca_excluir)
        st.write(
            f"Serão apagados: **{len(bancas_config.get(banca_excluir, []))}**"
            f" localidades, **{calendarios}** calendários e **{viagens}** viagens."
        )
        confirmar_banca = st.checkbox(
            "Confirmo a exclusão definitiva desta banca.", key="gestao_confirma_banca"
        )
        if st.button("🗑️ Excluir Banca", key="gestao_btn_rem_banca"):
            if not confirmar_banca:
                st.warning("Marque a confirmação antes de excluir.")
            else:
                salvar_snapshot(motivo="antes_remover_banca")
                for local in list(bancas_config.get(banca_excluir, [])):
                    _limpar_dados_da_localidade(banca_excluir, local)
                bancas_config.pop(banca_excluir, None)
                st.session_state["capacidade_bancas"].pop(banca_excluir, None)
                st.session_state["limite_fora_sede_bancas"].pop(banca_excluir, None)
                salvar_bancas(bancas_config)
                salvar_config("capacidade_bancas", st.session_state["capacidade_bancas"])
                salvar_config(
                    "limite_fora_sede_bancas",
                    st.session_state["limite_fora_sede_bancas"],
                )
                st.warning(f"Banca **{banca_excluir}** excluída.")
                st.rerun()


# ===========================================================================
# 7. APLICAÇÃO
# ===========================================================================
def configurar_logging() -> None:
    nivel = os.environ.get("DETRAN_LOG_LEVEL", "INFO").upper()
    logging.basicConfig(
        level=getattr(logging, nivel, logging.INFO),
        format="%(asctime)s | %(levelname)-8s | %(name)s | %(message)s",
    )


def painel_backup() -> None:
    with st.sidebar.expander("💾 Backup e Sincronização", expanded=False):
        st.caption(
            "Os dados ficam em banco SQLite e são gravados assim que você altera"
            " algo. Use os botões abaixo para guardar uma cópia externa ou"
            " restaurar um backup específico."
        )
        nome_arquivo = (
            f"backup_detran_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.json"
        )
        if st.button("📦 Preparar backup para download", key="btn_preparar_backup"):
            st.session_state["_backup_pronto"] = backup_em_json()

        if st.session_state.get("_backup_pronto"):
            st.download_button(
                "⬇️ Baixar backup",
                data=st.session_state["_backup_pronto"],
                file_name=nome_arquivo,
                mime="application/json",
                key="btn_baixar_backup",
            )

        if st.button("🗂️ Gravar snapshot no servidor", key="btn_snapshot"):
            destino = salvar_snapshot(motivo="manual")
            if destino:
                st.success(f"Snapshot gravado: `{destino.name}`")

        st.markdown("---")
        arquivo = st.file_uploader(
            "Restaurar um backup (.json):", type=["json"], key="upload_backup"
        )
        if arquivo is not None:
            confirmar = st.checkbox(
                "Confirmo que este backup deve substituir os dados atuais.",
                key="confirma_restauracao",
            )
            if st.button("♻️ Restaurar este backup", key="btn_restaurar"):
                if not confirmar:
                    st.warning("Marque a confirmação antes de restaurar.")
                else:
                    try:
                        aplicar_backup(json.load(arquivo))
                    except BackupInvalido as erro:
                        st.error(f"Backup recusado: {erro}")
                    except json.JSONDecodeError as erro:
                        st.error(f"O arquivo não é um JSON válido: {erro}")
                    except Exception as erro:
                        LOG.exception("Falha inesperada ao restaurar backup")
                        st.error(f"Falha ao restaurar: {erro}")
                    else:
                        st.success("Backup restaurado com sucesso!")
                        st.rerun()

        st.markdown("---")
        st.caption(
            "Se outra pessoa estiver usando o sistema ao mesmo tempo, recarregue"
            " para ver as alterações dela."
        )
        if st.button("🔄 Recarregar dados do banco", key="btn_recarregar"):
            recarregar_do_banco()
            st.rerun()


def barra_de_status() -> None:
    erro = st.session_state.get("erro_persistencia")
    if erro:
        st.error(
            "⚠️ O salvamento apresentou erro e os dados podem não ter sido"
            f" gravados: {erro}"
        )
    elif st.session_state.get("ultimo_salvamento"):
        st.caption(f"💾 Última gravação: {st.session_state['ultimo_salvamento']}")


def main() -> None:
    configurar_logging()

    st.set_page_config(
        page_title="DETRAN/MA - Gestão de Exames Práticos",
        page_icon="🚗",
        layout="wide",
        initial_sidebar_state="expanded",
    )
    st.markdown('<meta name="google" content="notranslate">', unsafe_allow_html=True)
    st.markdown(CSS_CUSTOMIZADO, unsafe_allow_html=True)

    carregar_estado()

    st.markdown(CABECALHO_HTML, unsafe_allow_html=True)
    painel_backup()
    barra_de_status()

    guia_quadro, guia_calendario, guia_viagens, guia_horarios, guia_relatorio, guia_gestao = st.tabs(
        [
            "🗓️ Quadro Geral de Examinadores",
            "📅 Montar Calendário Detalhado",
            "🚍 Gestão de Viagens Itinerantes",
            "⏰ Horários & Turmas",
            "📊 Dashboard Consolidado",
            "➕ Gestão de Localidades",
        ]
    )

    with guia_quadro:
        aba_quadro()
    with guia_calendario:
        aba_calendario()
    with guia_viagens:
        aba_viagens()
    with guia_horarios:
        aba_horarios()
    with guia_relatorio:
        aba_relatorio()
    with guia_gestao:
        aba_gestao()


if __name__ == "__main__":
    main()
